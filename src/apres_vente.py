"""Opérations fixes (après-vente) : bons de travail, heures, pièces.

Lit, pour chaque fichier source, le détail des opérations fixes :
  - par type de bon de travail (BT) : nombre de BT, main-d'œuvre (ventes et
    profit brut), pièces vendues sur ces BT (ventes et profit brut), heures
    vendues quand la source les donne ;
  - service : sous-traitance, fournitures, ajustements ;
  - pièces : comptoir, gros, pneus, huiles, ajustements, total ;
  - carrosserie ;
  - atelier : heures disponibles, pointées, vendues ; techniciens ; taux
    affichés.

Sources et ce qu'elles contiennent :
  - Réalisé (gabarit) : bloc Service (« M/O Client/Garantie/Interne/
    Esthétique », colonne # = nombre de BT), bloc Pièces, bloc Carrosserie.
    Budget et an passé inclus. Pas d'heures.
  - État GM (HAWKS, STM) : Page6 (nombre de B.R., ventes, profit brut par
    compte), Page4 (taux affichés), Page7 (techniciens). Pas d'heures.
  - État Hyundai Canada : Page 4 (# B.R., ventes, profit brut par compte),
    Page 6 (temps disponible, heures poinçonnées, temps facturé par type).
  - État Volkswagen Canada : Page 5 (B.R., ventes, profit brut par ligne),
    Page 6 (heures vendues par type, heures disponibles et productives).

Chaque montant est un dict {"real": x} (plus "budget" et "prior_year" pour le
Réalisé), comme le reste de data.json. Montants = ventes (M-O, pièces) ou
profit brut ; nombre de BT = compte de bons de travail par type (un BT qui a
des lignes client et garantie compte dans les deux types).
"""
import re
import unicodedata

import openpyxl

TYPES = ("client", "garantie", "interne", "esthetique")
MEASURES = ("bt", "mo_ventes", "mo_pb", "pc_ventes", "pc_pb", "heures")


def _sa(s):
    if not isinstance(s, str):
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("’", "'").replace("œ", "oe").replace("Œ", "oe")
    return re.sub(r"\s+", " ", s.lower()).strip()


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


class Sheet:
    """Feuille chargée une fois en mémoire (lecture seule, rapide)."""

    def __init__(self, ws):
        self.rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        self.title = ws.title

    def v(self, r, c):
        if r < 1 or r > len(self.rows):
            return None
        row = self.rows[r - 1]
        return row[c - 1] if 0 < c <= len(row) else None

    def n(self, r, c):
        return _num(self.v(r, c))

    def s(self, r, c):
        return _sa(self.v(r, c))

    @property
    def max_row(self):
        return len(self.rows)

    @property
    def max_col(self):
        return max((len(r) for r in self.rows), default=0)


def _kv(real=None, budget=None, prior=None):
    out = {"real": real}
    if budget is not None:
        out["budget"] = budget
    if prior is not None:
        out["prior_year"] = prior
    return out


def _add_kv(a, b):
    """Somme de deux {real, budget, prior_year} (None + x = x)."""
    if a is None:
        return dict(b) if b else None
    if b is None:
        return dict(a)
    out = {}
    for k in set(a) | set(b):
        x, y = a.get(k), b.get(k)
        out[k] = (x or 0) + (y or 0) if (x is not None or y is not None) else None
    return out


def _put(d, key, kv):
    d[key] = _add_kv(d.get(key), kv)


def _empty_section():
    return {"types": {}, "detail": {}, "service": {}, "pieces": {}, "carrosserie": {},
            "atelier": {}}


def _t(sec, typ):
    return sec["types"].setdefault(typ, {})


def _clean(sec):
    """Retire les types vides (aucun BT, aucun montant)."""
    def has(d):
        return any(isinstance(kv, dict) and any(v not in (None, 0) for v in kv.values())
                   for kv in d.values())
    for grp in ("types", "detail"):
        sec[grp] = {k: v for k, v in sec[grp].items() if has(v)}
    for grp in ("service", "pieces", "carrosserie"):
        sec[grp] = {k: v for k, v in sec[grp].items()
                    if (isinstance(v, dict) and has(v))}
    sec["atelier"] = {k: v for k, v in sec["atelier"].items() if v is not None}
    return sec


# ---------------------------------------------------------------------------
# Réalisé (gabarit)
# ---------------------------------------------------------------------------

# colonnes : (budget, réel, an passé)
G_CNT = (4, 10, 15)
G_SALES_SVC = (5, 11, 16)
G_PB = (6, 12, 17)
G_SALES_PC = (4, 10, 15)

G_SERVICE = {
    "m/o client": ("type", "client"),
    "m/o garantie": ("type", "garantie"),
    "m/o interne": ("type", "interne"),
    "m/o esthetique": ("type", "esthetique"),
    "ajustement m/o": ("service", "ajustement_mo"),
    "sous-total m/o": ("service", "total_mo"),
    "sous-traitement": ("service", "sous_traitance"),
    "total mecanique": ("service", "total"),
}
G_CARR = {
    "m/o carrosserie": "client",
    "carrosserie clients": "client",
    "carrosserie interne": "interne",
    "carrosserie concessions": "concessions",
    "m/o sous-traitement": "sous_traitance",
    "m/o esthetique carr.": "esthetique",
    "ajustement m/o carrosserie": "ajustement_mo",
    "total mecanique": "total",
    "total carrosserie": "total",
}
G_PIECES = {
    "pieces ro clients": ("type", "client"),
    "pieces ro carrosserie": ("pieces", "carrosserie"),
    "p/a garantie": ("type", "garantie"),
    "ventes internes": ("type", "interne"),
    "ventes details": ("pieces", "comptoir"),
    "ventes de gros": ("pieces", "gros"),
    "escomptes": ("pieces", "ajustements"),
    "rectif. inventaire": ("pieces", "ajustements"),
    "ventes pneus": ("pieces", "pneus"),
    "huiles et graisse": ("pieces", "huile"),
    "total dept. pieces": ("pieces", "total"),
}


def _g3(sh, r, cols):
    b, re_, p = (sh.n(r, c) for c in cols)
    return _kv(re_, b, p)


def gabarit_section(sh):
    sec = _empty_section()
    seen = set()
    for r in range(1, sh.max_row + 1):
        dept = sh.s(r, 2)
        label = sh.s(r, 3)
        if not label:
            continue
        key = (dept, label)
        if dept == "service" and label in G_SERVICE:
            if key in seen:
                continue
            seen.add(key)
            grp, name = G_SERVICE[label]
            cnt, sales, pb = _g3(sh, r, G_CNT), _g3(sh, r, G_SALES_SVC), _g3(sh, r, G_PB)
            if grp == "type":
                t = _t(sec, name)
                t["bt"], t["mo_ventes"], t["mo_pb"] = cnt, sales, pb
            else:
                sec["service"][name] = {"ventes": sales, "pb": pb}
                if name == "total_mo":
                    sec["service"][name]["bt"] = cnt
        elif dept == "carrosserie" and label in G_CARR:
            if key in seen:
                continue
            seen.add(key)
            name = G_CARR[label]
            sec["carrosserie"][name] = {"bt": _g3(sh, r, G_CNT), "ventes": _g3(sh, r, G_SALES_SVC),
                                        "pb": _g3(sh, r, G_PB)}
        elif dept == "pieces" and label in G_PIECES:
            if key in seen and G_PIECES[label][1] != "ajustements":
                continue
            seen.add(key)
            grp, name = G_PIECES[label]
            sales, pb = _g3(sh, r, G_SALES_PC), _g3(sh, r, G_PB)
            if grp == "type":
                t = _t(sec, name)
                t["pc_ventes"], t["pc_pb"] = sales, pb
            else:
                cur = sec["pieces"].setdefault(name, {})
                _put(cur, "ventes", sales if name != "ajustements" else None)
                _put(cur, "pb", pb)
    return _clean(sec)


def read_gabarit(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for sec_key, names in (("month", ("Réalisé Mois", "Print_Mois")), ("ytd", ("Réalisé AAD", "Print_AAD"))):
        name = next((n for n in names if n in wb.sheetnames), None)
        if name:
            out[sec_key] = gabarit_section(Sheet(wb[name]))
    wb.close()
    return out


# ---------------------------------------------------------------------------
# État GM (HAWKS, STM) -- Page6 : compte en colonne 5 ; mois : # 6, ventes 7,
# profit brut 9 ; cumul : # 11, ventes 12, profit brut 14.
# ---------------------------------------------------------------------------

GM_COLS = {"month": (6, 7, 9), "ytd": (11, 12, 14)}
GM_ACCOUNTS = {
    # main-d'œuvre mécanique
    "460A": ("mo", "client"), "460B": ("mo", "client"), "460C": ("mo", "client"),
    "460E": ("mo", "client"), "460F": ("mo", "client"), "460N": ("mo", "client"),
    "461A": ("mo", "client"), "461B": ("mo", "client"), "461C": ("mo", "client"),
    "461E": ("mo", "client"), "461F": ("mo", "client"),
    "460D": ("mo", "client", "mobile"),
    "462": ("mo", "garantie"),
    "463": ("mo", "interne"),
    "464": ("mo", "interne", "inspection"),
    "665": ("service", "ajustement_mo"),
    "469": ("service", "fournitures"),
    "466": ("service", "sous_traitance"),
    # carrosserie
    "470": ("carr", "client"), "471": ("carr", "client"),
    "472": ("carr", "garantie"), "473": ("carr", "interne"),
    "675": ("carr", "ajustement_mo"), "476": ("carr", "sous_traitance"),
    "479": ("carr", "peinture"),
    # pièces
    "480": ("pc", "garantie"),
    "467": ("pc", "client"), "467E": ("pc", "client"), "468": ("pc", "client"),
    "468E": ("pc", "client"), "478": ("pc", "client"), "467N": ("pc", "client"),
    "467D": ("pc", "client", "mobile"),
    "477": ("pieces", "carrosserie"),
    "481": ("pc", "interne"),
    "482": ("pieces", "comptoir"), "482C": ("pieces", "comptoir"),
    "483": ("pieces", "gros"),
    "484": ("pieces", "accessoires"), "484C": ("pieces", "accessoires"),
    "687": ("pieces", "ajustements"), "688": ("pieces", "ajustements"),
    "490": ("pieces", "pneus"), "491": ("pieces", "huile"), "492": ("pieces", "divers"),
}


def _gm_code(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip().upper()


def gm_sections(p6, p4=None, p7=None):
    out = {}
    for sec_key, (cc, cs, cp) in GM_COLS.items():
        sec = _empty_section()
        total_rows = {}
        for r in range(1, p6.max_row + 1):
            label = p6.s(r, 4) or p6.s(r, 3)
            if label.startswith("total mecanique"):
                total_rows["mecanique"] = r
            elif label.startswith("total carrosserie"):
                total_rows["carrosserie"] = r
            elif label.startswith("total operations fixes"):
                total_rows["fixes"] = r
            elif label.startswith("total pieces et access") and "pieces" not in total_rows:
                total_rows["pieces_partiel"] = r
            code = _gm_code(p6.v(r, 5))
            if code not in GM_ACCOUNTS:
                continue
            spec = GM_ACCOUNTS[code]
            grp, name = spec[0], spec[1]
            sub = spec[2] if len(spec) > 2 else None
            cnt, sales, pb = _kv(p6.n(r, cc)), _kv(p6.n(r, cs)), _kv(p6.n(r, cp))
            if grp == "mo":
                t = _t(sec, name)
                _put(t, "bt", cnt)
                _put(t, "mo_ventes", sales)
                _put(t, "mo_pb", pb)
                if sub:
                    d = sec["detail"].setdefault(sub, {})
                    _put(d, "bt", cnt)
                    _put(d, "mo_ventes", sales)
                    _put(d, "mo_pb", pb)
            elif grp == "pc":
                t = _t(sec, name)
                _put(t, "pc_ventes", sales)
                _put(t, "pc_pb", pb)
                if sub:
                    d = sec["detail"].setdefault(sub, {})
                    _put(d, "pc_ventes", sales)
                    _put(d, "pc_pb", pb)
            elif grp == "service":
                cur = sec["service"].setdefault(name, {})
                _put(cur, "ventes", sales)
                _put(cur, "pb", pb)
            elif grp == "carr":
                cur = sec["carrosserie"].setdefault(name, {})
                _put(cur, "bt", cnt)
                _put(cur, "ventes", sales)
                _put(cur, "pb", pb)
            elif grp == "pieces":
                cur = sec["pieces"].setdefault(name, {})
                _put(cur, "ventes", sales)
                _put(cur, "pb", pb)
        r = total_rows.get("mecanique")
        if r:
            sec["service"]["total"] = {"bt": _kv(p6.n(r, cc)), "ventes": _kv(p6.n(r, cs)),
                                       "pb": _kv(p6.n(r, cp))}
        r = total_rows.get("carrosserie")
        if r:
            sec["carrosserie"]["total"] = {"bt": _kv(p6.n(r, cc)), "ventes": _kv(p6.n(r, cs)),
                                           "pb": _kv(p6.n(r, cp))}
        # Total pièces : dernière ligne « TOTAL PIÈCES ET ACCESS. » (après pneus/huile)
        last_pc = None
        for rr in range(1, p6.max_row + 1):
            if (p6.s(rr, 4) or "").startswith("total pieces et access"):
                last_pc = rr
        if last_pc:
            sec["pieces"]["total"] = {"ventes": _kv(p6.n(last_pc, cs)), "pb": _kv(p6.n(last_pc, cp))}
        r = total_rows.get("fixes")
        if r:
            absorption = p6.n(r, cc)
            if absorption is not None:
                sec["atelier"]["absorption_etat"] = absorption
        if p4 is not None:
            sec["atelier"]["taux_affiche"] = gm_posted_rates(p4)
        if p7 is not None:
            techs = gm_techs(p7)
            if techs:
                sec["atelier"]["techniciens"] = techs
        out[sec_key] = _clean(sec)
    return out


def gm_posted_rates(p4):
    rates = {}
    for r in range(1, p4.max_row + 1):
        lab = p4.s(r, 2)
        if lab == "main-d'oeuvre client":
            rates["client"] = p4.n(r, 4)
            if p4.s(r, 8) == "garantie":
                rates["garantie"] = p4.n(r, 12)
        elif lab.startswith("main-d'oeuvre contrat prolonge") and p4.s(r, 8) == "interne":
            rates["interne"] = p4.n(r, 12)
    return {k: v for k, v in rates.items() if v}


def gm_techs(p7):
    """Techniciens (mécanique, carrosserie), Page7 « PERSONNEL EMPLOYÉ »."""
    hdr = None
    for r in range(1, p7.max_row + 1):
        row_labels = {c: p7.s(r, c) for c in range(1, min(p7.max_col, 20) + 1)}
        if "mecanique" in row_labels.values() and "carrosserie" in row_labels.values():
            hdr = {v: c for c, v in row_labels.items() if v}
            continue
        if hdr and p7.s(r, 3) == "techniciens":
            meca = p7.n(r, hdr["mecanique"])
            carr = p7.n(r, hdr["carrosserie"])
            return {"mecanique": meca, "carrosserie": carr}
    return None


def read_gm(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    p6 = Sheet(wb["Page6"])
    p4 = Sheet(wb["Page4"]) if "Page4" in wb.sheetnames else None
    p7 = Sheet(wb["Page7"]) if "Page7" in wb.sheetnames else None
    wb.close()
    return gm_sections(p6, p4, p7)


# ---------------------------------------------------------------------------
# État Hyundai Canada -- Page 4 : n° de compte en colonne 50 ; mois : # 2,
# ventes 7, profit brut 17 (« ( » en 16) ; cumul : # 53, ventes 58, profit
# brut 68 (« ( » en 67). Page 6 : heures (libellé col. 49, mois 60, cumul 66).
# ---------------------------------------------------------------------------

HY_COLS = {"month": (2, 7, 17), "ytd": (53, 58, 68)}
HY_ACCOUNTS = {
    370: ("mo", "client"), 372: ("mo", "garantie"), 374: ("mo", "interne"),
    378: ("mo", "esthetique"),
    480: ("service", "ajustement_mo"), 382: ("service", "sous_traitance"),
    384: ("service", "lubrifiants"), 383: ("service", "stockage_pneus"),
    385: ("service", "autres"),
    354: ("pc", "client"), 356: ("pc", "garantie"), 358: ("pc", "interne"),
    355: ("pieces", "carrosserie"),
    350: ("pieces", "gros"), 352: ("pieces", "comptoir"), 353: ("pieces", "accessoires"),
    360: ("pieces", "divers"), 362: ("pieces", "pneus"),
    463: ("pieces", "ajustements"), 464: ("pieces", "ajustements"),
}


def _hy_n(sh, r, c):
    v = sh.n(r, c)
    if v is None:
        return None
    left = sh.v(r, c - 1)
    if isinstance(left, str) and left.strip() == "(":
        return -v
    return v


def hy_sections(p4, p6, month_only=False):
    out = {}
    for sec_key, (cc, cs, cp) in HY_COLS.items():
        if month_only and sec_key == "ytd":
            continue
        sec = _empty_section()
        for r in range(1, p4.max_row + 1):
            acct = p4.v(r, 50)
            label = p4.s(r, 34)
            if isinstance(acct, (int, float)) and not isinstance(acct, bool) and int(acct) in HY_ACCOUNTS:
                grp, name = HY_ACCOUNTS[int(acct)]
                cnt, sales, pb = _kv(p4.n(r, cc)), _kv(p4.n(r, cs)), _kv(_hy_n(p4, r, cp))
                if grp == "mo":
                    t = _t(sec, name)
                    _put(t, "bt", cnt)
                    _put(t, "mo_ventes", sales)
                    _put(t, "mo_pb", pb)
                elif grp == "pc":
                    t = _t(sec, name)
                    _put(t, "pc_ventes", sales)
                    _put(t, "pc_pb", pb)
                elif grp == "service":
                    cur = sec["service"].setdefault(name, {})
                    _put(cur, "ventes", sales)
                    _put(cur, "pb", pb)
                elif grp == "pieces":
                    cur = sec["pieces"].setdefault(name, {})
                    _put(cur, "ventes", sales)
                    _put(cur, "pb", pb)
            elif label.startswith("total depart. pieces"):
                sec["pieces"]["total"] = {"ventes": _kv(p4.n(r, cs)), "pb": _kv(_hy_n(p4, r, cp))}
            elif label.startswith("total depart. service"):
                sec["service"]["total"] = {"ventes": _kv(p4.n(r, cs)), "pb": _kv(_hy_n(p4, r, cp))}
            elif label.startswith("total main d'oeuvre"):
                sec["service"]["total_mo"] = {"bt": _kv(p4.n(r, cc)), "ventes": _kv(p4.n(r, cs)),
                                              "pb": _kv(_hy_n(p4, r, cp))}
        if p6 is not None:
            hc = 60 if sec_key == "month" else 66
            at = sec["atelier"]
            for r in range(1, p6.max_row + 1):
                lab = p6.s(r, 49)
                if not lab:
                    continue
                val = p6.n(r, hc)
                if lab == "temps disponible":
                    at["heures_disponibles"] = val
                elif lab.startswith("heures poinconnees"):
                    at["heures_pointees"] = val
                elif lab == "temps facture":
                    at["heures_vendues"] = val
                elif lab == "temps facture - client":
                    _put(_t(sec, "client"), "heures", _kv(val))
                elif lab == "temps facture - garantie":
                    _put(_t(sec, "garantie"), "heures", _kv(val))
                elif lab == "temps facture - interne":
                    _put(_t(sec, "interne"), "heures", _kv(val))
                elif lab.startswith("temps facture - esthetique") or lab.startswith("temps facture - detaillage"):
                    if val:
                        _put(_t(sec, "esthetique"), "heures", _kv(val))
            for r in range(1, p6.max_row + 1):
                if p6.s(r, 2) == "techniciens":
                    t = p6.n(r, 17)
                    if t is not None:
                        at["techniciens"] = {"mecanique": t}
                    break
        out[sec_key] = _clean(sec)
    return out


def read_hyundai(path, month_only=False):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    p4 = Sheet(wb["Page 4"])
    p6 = Sheet(wb["Page 6"]) if "Page 6" in wb.sheetnames else None
    wb.close()
    return hy_sections(p4, p6, month_only=month_only)


# ---------------------------------------------------------------------------
# État Volkswagen Canada -- Page 5 (B.R., ventes, profit brut) et Page 6
# (heures). Trois générations de mise en page (français 2021-2023, français
# 2024-2025, anglais 2026) : colonnes repérées par les en-têtes.
# ---------------------------------------------------------------------------

def _vw_p5_columns(p5):
    """{'month': (cnt, sales, pb), 'ytd': (...)}, première ligne de données."""
    for r in range(1, min(p5.max_row, 12) + 1):
        labs = {c: p5.s(r, c) for c in range(1, p5.max_col + 1)}
        cnt = [c for c, v in labs.items() if v in ("b.r.", "ro's", "unites", "ro")]
        sales = [c for c, v in labs.items() if v in ("ventes", "sales")]
        pb = [c for c, v in labs.items() if v in ("profit brut", "gross profit")]
        if len(sales) >= 1 and len(pb) >= 1:
            cols = {}
            m_cnt = [c for c in cnt if c < sales[0]]
            cols["month"] = (m_cnt[-1] if m_cnt else None, sales[0], pb[0])
            if len(sales) >= 2 and len(pb) >= 2:
                y_cnt = [c for c in cnt if pb[0] < c < sales[1]]
                cols["ytd"] = (y_cnt[0] if y_cnt else None, sales[1], pb[1])
            return cols, r + 1
    raise ValueError("en-têtes de la Page 5 VW introuvables")


def _vw_label_col(p5, start):
    best = {}
    for r in range(start, min(p5.max_row, start + 30) + 1):
        for c in range(2, 8):
            v = p5.v(r, c)
            if isinstance(v, str) and len(v) > 12:
                best[c] = best.get(c, 0) + 1
    return max(best, key=best.get) if best else 3


def _vw_classify(label, section):
    """(groupe, nom) d'une ligne de la Page 5, ou None pour un sous-total."""
    s = label
    if not s or s.startswith(("total", "sub-total", "sous-total", "sous total", "grand total")):
        return None
    typ = None
    if any(k in s for k in ("garantie", "warranty")):
        typ = "garantie"
    elif any(k in s for k in ("interne", "internal")):
        typ = "interne"
    elif any(k in s for k in ("client", "customer", "prepaid", "prepay", "service rapide", "express")):
        typ = "client"
    if section == "service":
        if "non appliquee" in s or "unapplied" in s:
            return ("service", "ajustement_mo")
        if "sous-traitance" in s or "sublet" in s:
            return ("service", "sous_traitance")
        if "fournitures" in s or "shop supplies" in s:
            return ("service", "fournitures")
        if "esthetique" in s or "detailing" in s:
            return ("mo", "esthetique")
        if typ and ("main" in s or "labour" in s):
            return ("mo", typ)
        return None
    if section == "carrosserie":
        if "sous-traitance" in s or "sublet" in s:
            return ("carr", "sous_traitance")
        if "peinture" in s or "paint, compound" in s:
            return ("carr", "peinture")
        if typ:
            return ("carr", typ)
        return None
    if section == "pieces":
        if "ajustement" in s or "adjustment" in s:
            return ("pieces", "ajustements")
        if "pneus" in s or "tires" in s:
            return ("pieces", "pneus")
        if "comptoir" in s or "counter" in s:
            if "gros" in s or "wholesale" in s:
                return ("pieces", "gros")
            return ("pieces", "comptoir")
        if "carrosserie" in s or "p&b" in s or "peint & carr" in s:
            return ("pieces", "carrosserie")
        if typ:
            return ("pc", typ)
        return None
    return None


def vw_p5_sections(p5):
    cols, start = _vw_p5_columns(p5)
    lc = _vw_label_col(p5, start)
    out = {}
    for sec_key, (cc, cs, cp) in cols.items():
        sec = _empty_section()
        section = "service"  # la Page 5 commence par le service
        for r in range(start, p5.max_row + 1):
            marker = p5.s(r, 2)
            if marker.startswith("service"):
                section = "service"
            elif marker.startswith(("carrosserie", "paint & body", "peinture")):
                section = "carrosserie"
            elif marker.startswith(("pieces", "parts")):
                section = "pieces"
            label = p5.s(r, lc) or p5.s(r, lc + 1)
            if not label:
                continue
            cnt = _kv(p5.n(r, cc)) if cc else _kv(None)
            sales, pb = _kv(p5.n(r, cs)), _kv(p5.n(r, cp))
            if label.startswith(("total service", "total mecanique", "total mechanical (l")):
                if label.startswith("total service"):
                    sec["service"]["total"] = {"bt": cnt, "ventes": sales, "pb": pb}
                continue
            if label.startswith(("total peint", "total paint")):
                sec["carrosserie"]["total"] = {"bt": cnt, "ventes": sales, "pb": pb}
                continue
            if label.startswith(("total pieces et acc", "total parts & acc")):
                sec["pieces"]["total"] = {"ventes": sales, "pb": pb}
                continue
            cls = _vw_classify(label, section)
            if not cls:
                continue
            grp, name = cls
            if grp == "mo":
                t = _t(sec, name)
                _put(t, "bt", cnt)
                _put(t, "mo_ventes", sales)
                _put(t, "mo_pb", pb)
            elif grp == "pc":
                t = _t(sec, name)
                _put(t, "pc_ventes", sales)
                _put(t, "pc_pb", pb)
            elif grp == "service":
                cur = sec["service"].setdefault(name, {})
                _put(cur, "ventes", sales)
                _put(cur, "pb", pb)
            elif grp == "carr":
                cur = sec["carrosserie"].setdefault(name, {})
                _put(cur, "bt", cnt)
                _put(cur, "ventes", sales)
                _put(cur, "pb", pb)
            elif grp == "pieces":
                cur = sec["pieces"].setdefault(name, {})
                _put(cur, "ventes", sales)
                _put(cur, "pb", pb)
        out[sec_key] = sec
    return out


def _find(sh, pred, cols=None):
    for r in range(1, sh.max_row + 1):
        for c in (cols or range(1, sh.max_col + 1)):
            if pred(sh.s(r, c)):
                return r, c
    return None, None


def vw_p6_hours(p6, out):
    # Heures vendues par type
    r0, c0 = _find(p6, lambda s: s.startswith("heures vendues - service") or s.startswith("sold hours - service"))
    if r0:
        hdr = r0 + 1
        mcols, ycols = [], []
        for c in range(c0 + 1, c0 + 25):
            h = p6.s(hdr, c)
            if not h:
                continue
            if re.search(r"(^|\s)(mc|mtd|mois)(\s|$|-)", h):
                mcols.append(c)
            elif re.search(r"(^|\s)(ac|ytd|cumul)(\s|$|-)", h):
                ycols.append(c)
        typmap = {"client": "client", "customer": "client", "garantie": "garantie",
                  "warranty": "garantie", "interne": "interne", "internal": "interne",
                  "heures ppm": "client", "ppm hours": "client"}
        for r in range(r0 + 1, r0 + 8):
            lab = p6.s(r, c0)
            if lab in typmap:
                for sec_key, cs in (("month", mcols), ("ytd", ycols)):
                    if sec_key not in out:
                        continue
                    vals = [p6.n(r, c) for c in cs]
                    vals = [v for v in vals if v is not None]
                    if vals:
                        _put(_t(out[sec_key], typmap[lab]), "heures", _kv(sum(vals)))
    # Analyse d'atelier
    r1, c1 = _find(p6, lambda s: s.startswith(("tot. heures disp", "tot. avail")))
    if r1:
        hdr = None
        for rr in range(r1 - 1, r1 - 4, -1):
            if any(p6.s(rr, c).startswith(("serv", "service")) for c in range(c1, c1 + 25)):
                hdr = rr
                break
        mcol = ycol = None
        if hdr:
            for c in range(c1, c1 + 25):
                h = p6.s(hdr, c)
                if h.startswith(("serv", "service")):
                    if any(k in h for k in ("cumul", "ytd")):
                        ycol = ycol or c
                    else:
                        mcol = mcol or c
        rows = {}
        for rr in range(r1, r1 + 5):
            lab = p6.s(rr, c1)
            if lab.startswith(("tot. heures disp", "tot. avail")):
                rows["heures_disponibles"] = rr
            elif lab.startswith(("tot. heures prod", "tot. prod")):
                rows["heures_pointees"] = rr
            elif (lab.startswith("hres vendu") and "vw" in lab) or lab.startswith("vw sold"):
                rows["heures_vendues"] = rr
            elif (lab.startswith("hres vendu") and "autre" in lab) or lab.startswith("o.m. sold"):
                rows["heures_vendues_autres"] = rr
        for sec_key, col in (("month", mcol), ("ytd", ycol)):
            if sec_key not in out or not col:
                continue
            at = out[sec_key]["atelier"]
            for k, rr in rows.items():
                v = p6.n(rr, col)
                if v is not None:
                    at[k] = v
            if at.get("heures_vendues_autres"):
                at["heures_vendues"] = (at.get("heures_vendues") or 0) + at.pop("heures_vendues_autres")
            else:
                at.pop("heures_vendues_autres", None)
    # Techniciens
    techs = 0.0
    found = False
    for r in range(1, p6.max_row + 1):
        for c in range(30, min(p6.max_col, 60) + 1):
            s = p6.s(r, c)
            if s in ("vw techs", "vw technicians", "techniciens - vw", "techniciens - autre",
                     "technicians - other", "techniciens - autres"):
                for cc in range(c + 1, c + 14):
                    v = p6.n(r, cc)
                    if v is not None:
                        techs += v
                        found = True
                        break
    if found:
        for sec_key in out:
            out[sec_key]["atelier"]["techniciens"] = {"mecanique": techs}
    return out


def read_vw(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    p5 = Sheet(wb["Page 5"])
    p6 = Sheet(wb["Page 6"]) if "Page 6" in wb.sheetnames else None
    wb.close()
    out = vw_p5_sections(p5)
    if p6 is not None:
        vw_p6_hours(p6, out)
    return {k: _clean(v) for k, v in out.items()}


# ---------------------------------------------------------------------------

def read(path, fmt, sections=("month", "ytd")):
    """Opérations fixes d'un fichier, par section présente dans l'extraction."""
    if fmt == "etat_gm":
        out = read_gm(path)
    elif fmt == "etat_hyundai":
        out = read_hyundai(path, month_only="ytd" not in sections)
    elif fmt == "etat_vw":
        out = read_vw(path)
    else:
        out = read_gabarit(path)
    return {k: v for k, v in out.items() if k in sections}


def _bt_total(sec):
    tot = 0.0
    for t in (sec.get("types") or {}).values():
        tot += ((t.get("bt") or {}).get("real") or 0)
    return tot


def merge_state_hours(target, state, tol=0.03):
    """Le Réalisé n'a pas d'heures : on reprend celles de l'état du
    constructeur du même mois (VW, Hyundai) si les nombres de BT concordent
    (écart ≤ 3 %). Renvoie True si les heures ont été reprises."""
    if not target or not state:
        return False
    bt_t, bt_s = _bt_total(target), _bt_total(state)
    if not bt_t or not bt_s or abs(bt_t - bt_s) / bt_t > tol:
        return False
    moved = False
    for typ, meas in (state.get("types") or {}).items():
        h = (meas.get("heures") or {}).get("real")
        if h:
            target.setdefault("types", {}).setdefault(typ, {})["heures"] = {"real": h}
            moved = True
    for k, v in (state.get("atelier") or {}).items():
        target.setdefault("atelier", {}).setdefault(k, v)
    if moved:
        target["heures_source"] = "etat"
    return moved


def _stats_signature(sec):
    """Empreinte des statistiques non financières (BT et heures par type)."""
    if not sec:
        return None
    items = []
    for typ in sorted(sec.get("types") or {}):
        t = sec["types"][typ]
        items.append((typ, (t.get("bt") or {}).get("real"), (t.get("heures") or {}).get("real")))
    if not any(b for _, b, _ in items):
        return None
    return tuple(items)


def drop_duplicate_stats(store, name_period):
    """Un état dont les nombres de BT et les heures sont identiques, au BT près,
    à ceux d'un autre mois (ex. « FFS Hyundai F 09-2026 », qui contient les
    chiffres financiers de septembre 2025 mais les statistiques d'atelier
    d'août 2026) : les statistiques ne sont gardées que dans le mois dont le
    nom de fichier correspond au contenu ; ailleurs, BT et heures sont retirés
    (les montants restent). Renvoie la liste des (concession, mois, section)."""
    removed = []
    for dealer, dd in (store.get("dealers") or {}).items():
        for sk in ("month", "ytd"):
            groups = {}
            for p, rec in (dd.get("periods") or {}).items():
                av = ((rec.get("sections") or {}).get(sk) or {}).get("apres_vente")
                sig = _stats_signature(av)
                if sig:
                    groups.setdefault(sig, []).append(p)
            for sig, periods in groups.items():
                if len(periods) < 2:
                    continue
                owners = [p for p in periods
                          if name_period((dd["periods"][p].get("source_file") or "")) == p]
                keep = owners[0] if len(owners) == 1 else None
                for p in periods:
                    if p == keep:
                        continue
                    av = dd["periods"][p]["sections"][sk]["apres_vente"]
                    for grp in ("types", "detail"):
                        for t in (av.get(grp) or {}).values():
                            t.pop("bt", None)
                            t.pop("heures", None)
                    for k in ("heures_disponibles", "heures_pointees", "heures_vendues", "techniciens"):
                        (av.get("atelier") or {}).pop(k, None)
                    for k in ("service", "carrosserie"):
                        for v in (av.get(k) or {}).values():
                            if isinstance(v, dict):
                                v.pop("bt", None)
                    av["stats_retirees"] = f"BT et heures identiques à ceux de {', '.join(x for x in periods if x != p)}"
                    removed.append((dealer, p, sk))
    return removed


if __name__ == "__main__":
    import json
    import sys
    from extract import detect_format  # noqa: E402
    for p in sys.argv[1:]:
        fmt = detect_format(p)
        print(p, fmt)
        print(json.dumps(read(p, fmt), ensure_ascii=False, indent=1)[:6000])
