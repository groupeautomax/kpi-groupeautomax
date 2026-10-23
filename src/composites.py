"""Composites des constructeurs : moyennes du groupe de comparaison (Hyundai)
et bulletin de la concession (Volkswagen), lus dans les PDF que les
constructeurs publient chaque mois.

Formats lus :
  - Hyundai Canada, eComposite™ « ÉF ‐ Sommaire » : moyennes du groupe de
    comparaison (ex. « Est A (850+) (16 / 16) ») pour le mois (Moy. MAD) et
    le cumul (Moy. PAD), avec l'écart de ces moyennes par rapport à l'an
    passé. Ne contient PAS les chiffres de la concession : ceux-ci sont
    calculés à partir de l'état financier Hyundai Canada (FFS) du même mois,
    avec les mêmes définitions (voir hyundai_statement_lines).
  - Hyundai Canada, eComposite™ « P&P » (profits et pertes par département :
    Total des opérations, Neuf, Occasion, Pièces, Service, Carrosserie,
    Location et bail) : Moy. MAD, Moy. MMAP (même mois an passé), Moy. PAD,
    MPAD moy. (même période an passé).
  - Volkswagen Canada, « Dealer Report Card » : chiffres de la concession
    (2026 YTD, 2025 YTD, mois, 3 mois), moyennes nationale et Québec (Geo
    Local), objectifs, points et rangs.
  - Volkswagen Canada, états financiers standard d'un groupe de volume (Excel,
    « Financial Statement_Volume Group_A - ≥675_Standard statement ») : le
    concessionnaire moyen du groupe, au gabarit de l'état de la concession
    (voir vw_statement_lines, qui lit aussi l'état de VW Brossard).

Le PDF n'est jamais déposé dans le dépôt (public) : seuls les chiffres, en
JSON, vont dans sources/composites/ (voir drive_sync.py et write_json).

Demande de Maxime Allard (29 septembre 2026) : ajouter les comparaisons au
composite dans les rapports ; BMW et GM suivront.

    python src/composites.py fichier.pdf [...] [--sortie sources/composites]
"""

import datetime as _dt
import json
import os
import re
import sys
import unicodedata

MOIS_FR = {"janv": 1, "févr": 2, "fevr": 2, "mars": 3, "avr": 4, "mai": 5, "juin": 6, "juil": 7,
           "août": 8, "aout": 8, "sept": 9, "oct": 10, "nov": 11, "déc": 12, "dec": 12}
MONTHS_EN = {m: i for i, m in enumerate(
    ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}

# Version des lecteurs : drive_sync.py relit un PDF déjà refusé quand elle
# change (nouveau format pris en charge, ex. BMW ou GM).
PARSER_VERSION = 5   # 4 : exports Excel de l'eComposite Hyundai ; 5 : états standard du groupe de volume VW

PAREN_OPEN = "(\ufd3e"      # « ﴾ » dans les PDF eComposite
PAREN_CLOSE = ")\ufd3f"     # « ﴿ »


def norm(s):
    """Minuscules, sans accents, tirets et espaces normalisés."""
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("\u2010", "-").replace("\u2011", "-").replace("\u2013", "-").replace("\u2014", "-")
    s = s.replace("\ufd3e", "(").replace("\ufd3f", ")").replace("\u2019", "'")
    s = s.replace("œ", "oe").replace("Œ", "OE")
    return re.sub(r"\s+", " ", s.lower()).strip()


# --------------------------------------------------------------- nombres
NUM_RE = re.compile(r"^[" + PAREN_OPEN + r"]?-?[\d\s\u00a0\u202f]*\d(?:[.,]\d+)?\s?[$%]?[" + PAREN_CLOSE + r"]?$")


def parse_num(tok):
    """« 7 505 306 $ » → 7505306 ; « ﴾10,8 %﴿ » → -0.108 ; « 102 » → 102 ;
    « 0$ » → 0. Renvoie (valeur, unité) avec unité '$', '%' ou ''."""
    t = tok.strip()
    neg = t[:1] in PAREN_OPEN or t[-1:] in PAREN_CLOSE
    t = t.strip(PAREN_OPEN + PAREN_CLOSE).strip()
    unit = ""
    if t.endswith("$") or t.endswith("%"):
        unit, t = t[-1], t[:-1].strip()
    if t.startswith("-"):
        neg, t = not neg, t[1:]
    t = re.sub(r"[\s\u00a0\u202f]", "", t).replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return None, None
    if unit == "%":
        v /= 100
    return (-v if neg else v), unit


def is_num(tok):
    return bool(NUM_RE.match(tok.strip())) and parse_num(tok)[0] is not None


# ------------------------------------------------------------ lecture PDF
def pdf_lines(path, gap=8.0):
    """[(page, top, [(x0, texte)])] : mots regroupés en jetons (écart < gap
    points : un nombre « 7 505 306 $ » reste un seul jeton, deux colonnes
    sont séparées)."""
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            words = page.extract_words(x_tolerance=1.5, keep_blank_chars=False)
            rows = {}
            for w in words:
                rows.setdefault(round(w["top"] / 2.0), []).append(w)
            for top in sorted(rows):
                ws = sorted(rows[top], key=lambda w: w["x0"])
                toks, cur, prev = [], None, None
                for w in ws:
                    if prev is not None and w["x0"] - prev["x1"] < gap:
                        cur[1] += " " + w["text"]
                    else:
                        cur = [w["x0"], w["text"]]
                        toks.append(cur)
                    prev = w
                out.append((pno, top * 2.0, [(x, t) for x, t in toks]))
    return out


def pdf_text(path, max_pages=None):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        pages = pdf.pages[:max_pages] if max_pages else pdf.pages
        return "\n".join(p.extract_text() or "" for p in pages)


# ------------------------------------------------------ eComposite Hyundai
HY_SKIP = ("de:", "donnees:", "% d'ecart", "mad", "pad", "description", "departement", "econposite", "ecomposite")
# Départements du « ÉF ‐ Sommaire » (colonne de gauche)
HY_EF_DEPTS = {"totaux": "total", "vehicule neuf": "neuf", "vehicules d'occasion": "occasion", "service": "service",
               "pieces": "pieces", "carrosserie": "carrosserie", "bail et location": "location"}
# Titres des pages P&P
HY_PP_PAGES = {"total des operations": "total", "neuf": "neuf", "occasion": "occasion", "pieces": "pieces",
               "service": "service", "carrosserie": "carrosserie", "location et bail": "location"}
HY_PP_BLOCKS = {"les frais de vente de vehicules": "vente", "frais directs": "directs", "frais fixes": "fixes"}


def _period_fr(text):
    """« De: janv.‐2026 À: août‐2026 » → ((2026, 1), (2026, 8))."""
    t = norm(text)
    m = re.search(r"de:\s*([a-z]+)\.?\s*-?\s*(20\d\d)\s*a:\s*([a-z]+)\.?\s*-?\s*(20\d\d)", t)
    if not m:
        return None, None
    a = MOIS_FR.get(m.group(1)[:4]) or MOIS_FR.get(m.group(1)[:3])
    b = MOIS_FR.get(m.group(3)[:4]) or MOIS_FR.get(m.group(3)[:3])
    if not a or not b:
        return None, None
    return (int(m.group(2)), a), (int(m.group(4)), b)


def _group(text, lines=None):
    """« Est A ﴾850+﴿ ﴾16 / 16﴿ » → ("Est A (850+)", 16, 16). Avec lines, on
    cherche d'abord un jeton qui ne contient que le groupe (le titre de
    section peut être sur la même ligne : « Voitures - Neuf »)."""
    for _, _, toks in (lines or []):
        for _, tok in toks:
            if re.search(r"\(\s*\d+\s*/\s*\d+\s*\)", tok.replace("\ufd3e", "(").replace("\ufd3f", ")")):
                g = _group(tok)
                if g[0]:
                    return g
    t = text.replace("\ufd3e", "(").replace("\ufd3f", ")")
    m = re.search(r"([A-Za-zÀ-ÿ][^\n()]*?\(\s*\d+\s*\+?\s*\))\s*\(\s*(\d+)\s*/\s*(\d+)\s*\)", t)
    if not m:
        return None, None, None
    return re.sub(r"\s+", " ", m.group(1)).strip(), int(m.group(2)), int(m.group(3))


def _rows(lines, ncols, label_x=None):
    """Lignes de données : [(page, étiquette, [valeurs])] ; une étiquette sur
    plusieurs lignes (texte seul avant et/ou après la ligne de chiffres) est
    recollée. Les jetons à gauche de label_x (colonne « Département ») sont
    rendus à part : [(page, 'dept', texte)]."""
    items = []
    for pno, top, toks in lines:
        vals = [t for x, t in toks if is_num(t)]
        texts = [(x, t) for x, t in toks if not is_num(t)]
        if len(vals) >= ncols:
            vals = vals[-ncols:]
            lab = " ".join(t for x, t in texts if label_x is None or x >= label_x - 10)
            dept = [t for x, t in texts if label_x is not None and x < label_x - 10]
            for d in dept:
                items.append(("dept", pno, d))
            items.append(("row", pno, lab, [parse_num(v) for v in vals]))
        elif not vals and texts:
            for x, t in texts:
                if label_x is not None and x < label_x - 10:
                    items.append(("dept", pno, t))
                else:
                    items.append(("text", pno, t))
    # recoller les étiquettes coupées
    out = []
    i = 0
    while i < len(items):
        it = items[i]
        if it[0] == "row":
            lab = it[2]
            if not lab and out and out[-1][0] == "text":
                lab = out.pop()[2]
            j = i + 1
            # suite de l'étiquette : texte seul qui suit, sauf s'il annonce la ligne suivante
            while j < len(items) and items[j][0] == "text" and not _is_header(items[j][2]):
                nxt = items[j + 1] if j + 1 < len(items) else None
                if nxt and nxt[0] == "row" and not nxt[2]:
                    break
                lab = (lab + " " + items[j][2]).strip()
                j += 1
            out.append(("row", it[1], lab, it[3]))
            i = j
            continue
        out.append(it)
        i += 1
    return out


def _is_header(t):
    n = norm(t)
    return any(n.startswith(h) for h in HY_SKIP) or "ecomposite" in n or re.fullmatch(r"page \d+", n) is not None


def _label_x(lines, word="Description"):
    for pno, top, toks in lines:
        for x, t in toks:
            if t.strip().startswith(word):
                return x
    return None


def read_hyundai_ef(path, lines=None):
    """« ÉF ‐ Sommaire » : {"kind": "hyundai_ef", "period", "groupe", "n",
    "lignes": {dept: {id: {"mad", "mad_ecart", "pad", "pad_ecart"}}}}."""
    lines = lines or pdf_lines(path)
    text = "\n".join(" ".join(t for _, t in toks) for _, _, toks in lines)
    (y0, m0), (y1, m1) = _period_fr(text)
    grp, n, n_tot = _group(text, lines)
    lx = _label_x(lines)
    dept, out = None, {}
    for it in _rows(lines, 4, lx):
        if it[0] == "dept":
            d = HY_EF_DEPTS.get(norm(it[2]))
            if d:
                dept = d
            continue
        if it[0] != "row":
            continue
        lab = it[2]
        (mad, u1), (e1, _), (pad, u2), (e2, _) = it[3]
        out.setdefault(dept or "?", []).append({"label": re.sub(r"\s+", " ", lab).strip(), "unit": u1 or u2,
                                                 "mad": mad, "mad_ecart": e1, "pad": pad, "pad_ecart": e2})
    # Le département est écrit au milieu de son bloc : on réaffecte par
    # l'ordre des lignes (chaque bloc commence par une ligne connue).
    rows = [r for rs in out.values() for r in rs]
    return {"kind": "hyundai_ef", "source": "eComposite Hyundai Canada — ÉF ‐ Sommaire",
            "period": f"{y1}-{m1:02d}", "debut": f"{y0}-{m0:02d}", "groupe": grp, "n": n, "n_total": n_tot,
            "lignes": ef_assign(rows)}


# Ordre des lignes du « ÉF ‐ Sommaire » : (département, identifiant, début de l'étiquette)
HY_EF_LINES = [
    ("total", "ventes", "total des ventes"), ("total", "pb", "total du profit brut"), ("total", "frais", "frais total $"),
    ("total", "pn_avant_bonis", "profit net avant bonis"), ("total", "pn_pct_pb", "benefice net avant primes"),
    ("total", "publicite", "frais total de publicite"), ("total", "absorption", "absorption"),
    ("neuf", "unites", "total detail neuf"), ("neuf", "ventes", "neuf - ventes"), ("neuf", "pb", "pb - neuf"),
    ("neuf", "frais", "total des depenses"), ("neuf", "profit_op", "profit d'operation net"),
    ("neuf", "profit_op_pct_pb", "profit d'operation %"), ("neuf", "fi", "bureau commercial pb - neuf"),
    ("neuf", "fi_unite", "bureau commercial pb pvnd"),
    ("occasion", "unites", "total detail d'occasion"), ("occasion", "ventes", "ventes - occasion"),
    ("occasion", "pb", "pb - occasion"), ("occasion", "frais", "total des depenses"),
    ("occasion", "profit_op", "profit d'operation net"), ("occasion", "profit_op_pct_pb", "profit d'operation %"),
    ("occasion", "fi", "bureau commercial pb - occasion"), ("occasion", "fi_unite", "bureau commercial pb pvod"),
    ("service", "ventes", "ventes de service"), ("service", "pb", "pb - service"), ("service", "frais", "total des depenses"),
    ("service", "profit_op", "profit d'operation net"), ("service", "profit_op_pct_pb", "profit d'operation %"),
    ("pieces", "ventes", "ventes - pieces"), ("pieces", "pb", "pb - pieces"), ("pieces", "frais", "total des depenses"),
    ("pieces", "profit_op", "profit d'operation net"), ("pieces", "profit_op_pct_pb", "profit d'operation %"),
    ("carrosserie", "ventes", "ventes - carrosserie"), ("carrosserie", "pb", "pb - carrosserie"),
    ("carrosserie", "frais", "total des depenses"), ("carrosserie", "profit_op", "profit d'operation net"),
    ("carrosserie", "profit_op_pct_pb", "profit d'operation %"),
    ("location", "ventes", "ventes - bail"), ("location", "pb", "pb - bail"), ("location", "frais", "frais total"),
    ("location", "profit_op", "profit d'operation net"), ("location", "profit_op_pct_pb", "profit d'operation %"),
]


def ef_assign(rows):
    """Associe les lignes lues à HY_EF_LINES, dans l'ordre ; une ligne
    inattendue arrête la lecture (mise en page changée)."""
    out, i = {}, 0
    for r in rows:
        while i < len(HY_EF_LINES) and not norm(r["label"]).startswith(HY_EF_LINES[i][2]):
            i += 1
        if i >= len(HY_EF_LINES):
            raise ValueError(f"ÉF ‐ Sommaire : ligne inattendue « {r['label']} »")
        dept, key, _ = HY_EF_LINES[i]
        out.setdefault(dept, {})[key] = {k: r[k] for k in ("mad", "mad_ecart", "pad", "pad_ecart")}
        i += 1
    missing = [f"{d}.{k}" for d, k, _ in HY_EF_LINES if k not in out.get(d, {})]
    if len(missing) > 10:
        raise ValueError(f"ÉF ‐ Sommaire incomplet : {', '.join(missing[:6])}…")
    return out


def read_hyundai_pp(path, lines=None):
    """« P&P » : {"kind": "hyundai_pp", "period", "groupe", "n",
    "pages": {dept: [{"bloc", "label", "mad", "mmap", "pad", "mpad"}]}}."""
    lines = lines or pdf_lines(path)
    text = "\n".join(" ".join(t for _, t in toks) for _, _, toks in lines)
    (y0, m0), (y1, m1) = _period_fr(text)
    grp, n, n_tot = _group(text, lines)
    pages = {}
    dept, bloc = None, None
    for it in _rows(lines, 6):
        if it[0] == "text":
            t = norm(it[2])
            t = re.sub(r"\s*-?\s*p&p$", "", t).strip()
            if t in HY_PP_PAGES:
                dept = HY_PP_PAGES[t]
            b = next((v for k, v in HY_PP_BLOCKS.items() if t.startswith(k)), None)
            if b:
                bloc = b
            continue
        if it[0] != "row" or not dept:
            continue
        lab = re.sub(r"\s+", " ", it[2]).strip()
        for k in HY_PP_BLOCKS:     # « Frais directs Est A (850+) » sur la même ligne que le groupe
            if norm(lab).startswith(k):
                bloc = HY_PP_BLOCKS[k]
        (mad, u), (mmap, _), _e1, (pad, _), (mpad, _), _e2 = it[3]
        pages.setdefault(dept, []).append({"bloc": bloc, "label": lab, "unit": u,
                                           "mad": mad, "mmap": mmap, "pad": pad, "mpad": mpad})
    if not pages.get("total"):
        raise ValueError("P&P : page « Total des opérations » introuvable")
    return {"kind": "hyundai_pp", "source": "eComposite Hyundai Canada — P&P",
            "period": f"{y1}-{m1:02d}", "debut": f"{y0}-{m0:02d}", "groupe": grp, "n": n, "n_total": n_tot,
            "pages": pages}


AR_SUB_RE = re.compile(r"^(ventes( par | pv|$)|pb( |%|$)|\$pb|penetration|remise en etat)")
# Postes de l'analyse des revenus comparés à l'état FFS (Page 4) : (page, id, début
# de l'étiquette du composite, début de l'étiquette de l'état, ligne d'en-tête de l'état)
AR_LINES = [
    ("Service ‐ Vente et PB", "mo_client", "main-d'oeuvre - bons de reparation au client", "main-d'oeuvre - bons de reparation", "service"),
    ("Service ‐ Vente et PB", "mo_garantie", "main-d'oeuvre - garantie", "main d'oeuvre - garantie", "service"),
    ("Service ‐ Vente et PB", "mo_interne", "main-d'oeuvre - interne", "main d'oeuvre - interne", "service"),
    ("Service ‐ Vente et PB", "mo_total", "sous total main-d'oeuvre", "total main d'oeuvre", "service"),
    ("Service ‐ Vente et PB", "sous_traitance", "reparation confiees aux sous-traitants", "reparations confiees", "service"),
    ("Service ‐ Vente et PB", "lubrifiants", "lubrifiants", "lubrifiants", "service"),
    ("Pièces ‐ Vente et PB", "pc_gros", "ventes en gros", "ventes en gros", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_comptoir", "ventes au comptoir, pieces", "ventes au comptoir - pieces", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_accessoires", "ventes au comptoir - accessoires", "ventes au comptoir - accessoires", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_client", "bons de reparation au client", "bons de reparation au client, service", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_garantie", "garantie", "garantie", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_interne", "interne", "interne", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_pneus", "roues & pneus", "roues & pneus", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_boni_gros", "escompte / boni pour ventes en gros", "escompte / boni", "pieces"),
    ("Pièces ‐ Vente et PB", "pc_total", "total depart. pieces", "total depart. pieces", "pieces"),
]


def ar_item(ar, page, prefix):
    """Poste de l'analyse des revenus : {"n": {mad…}, "ventes": {…}, "pb": {…}}."""
    rows = (ar.get("pages") or {}).get(page) or []
    for it in rows:
        if norm(it["label"]).replace("‐", "-").startswith(prefix):
            sous = {norm(k): v for k, v in (it.get("sous") or {}).items()}
            return {"n": it.get("tete"), "ventes": sous.get("ventes"), "pb": sous.get("pb")}
    return None


def read_hyundai_ar(path, lines=None):
    """« Analyse des revenus » (eComposite) : par page (Neuf ‐ Ventes, PB ‐ Neuf,
    F et A Neuf, F et A Occasion, Occasion ‐ Ventes, Occasion ‐ PB, Pièces ‐
    Vente et PB, Service ‐ Vente et PB, Carrosserie, Location et bail, Autres
    revenus, Autres déductions), des postes : ligne de tête (nombre de
    transactions ou d'unités, parfois un montant) puis sous-lignes (Ventes,
    Ventes par transaction, PB, PB % des ventes…). Colonnes : Moy. MAD, Moy.
    MMAP, Moy. PAD, MPAD moy. (écarts non gardés)."""
    lines = lines or pdf_lines(path)
    text = "\n".join(" ".join(t for _, t in toks) for _, _, toks in lines)
    (y0, m0), (y1, m1) = _period_fr(text)
    grp, n, n_tot = _group(text, lines)
    pages, cur_page, cur = {}, None, None
    # colonne des postes de tête : la plus à gauche des lignes chiffrées de la page
    head_x = {}
    for pno, top, toks in lines:
        if sum(1 for _, t in toks if is_num(t)) >= 4:
            head_x[pno] = min(head_x.get(pno, 1e9), toks[0][0])
    for pno, top, toks in lines:
        texts = [(x, t) for x, t in toks if not is_num(t) and t.strip() not in ("N/A",)]
        vals = [t for x, t in toks if is_num(t)]
        joined = " ".join(t for _, t in texts)
        m = re.search(r"analyse des revenus\s*>>\s*(.+?)\s*$", norm(joined).replace("\\", ""))
        if m and "ecomposite" in norm(joined):
            continue          # pied de page (titre repris en tête de page suivante)
        if not texts and not vals:
            continue
        x0 = toks[0][0]
        if not vals:
            t = norm(joined)
            if _is_header(joined) or t.startswith("donnees") or "est a" in t or t.startswith("% d'ecart"):
                continue
            # titre de page (grand, à gauche) : « Service ‐ Vente et PB »
            if top < 45:
                title = re.sub(r"\s+", " ", joined).strip()
                if title != cur_page:          # même titre = suite de la page précédente
                    cur_page, cur = title, None
                    pages.setdefault(cur_page, [])
                continue
            # poste sans valeurs (« Temps non‐assigné », « Divers autos ») ou groupe (« Bureau commercial ‐ Neuf »)
            if cur_page:
                cur = {"label": re.sub(r"\s+", " ", joined).strip(), "tete": None, "sous": {}}
                pages[cur_page].append(cur)
            continue
        if not cur_page or len(vals) < 4:
            continue
        v = [parse_num(t)[0] for t in vals]
        if len(v) >= 6:
            q = {"mad": v[0], "mmap": v[1], "pad": v[3], "mpad": v[4]}
        else:                   # « N/A » : colonnes an passé absentes
            q = {"mad": v[0], "mmap": None, "pad": v[2] if len(v) > 2 else None, "mpad": None}
        lab = re.sub(r"\s+", " ", joined).strip()
        # sous-ligne : plus à droite que les postes, ou libellé de sous-ligne en haut
        # d'une page de suite (« Ventes », « PB », « PB % des ventes »…)
        is_sub = cur is not None and (x0 > head_x.get(pno, 0) + 3 or AR_SUB_RE.match(norm(lab)) is not None)
        if not is_sub:
            cur = {"label": lab, "tete": q, "sous": {}}
            pages[cur_page].append(cur)
        else:
            cur["sous"][lab] = q
    # postes vides (groupes) retirés
    pages = {k: [it for it in v if it["tete"] or it["sous"]] for k, v in pages.items()}
    if not any(pages.values()):
        raise ValueError("Analyse des revenus : aucun poste lu")
    return {"kind": "hyundai_ar", "source": "eComposite Hyundai Canada — Analyse des revenus",
            "period": f"{y1}-{m1:02d}", "debut": f"{y0}-{m0:02d}", "groupe": grp, "n": n, "n_total": n_tot,
            "pages": pages}


# ------------------------------------------------ eComposite Hyundai (Excel)
# Les mêmes rapports que les PDF, exportés en Excel par Mustafa Bhaloo (22 octobre
# 2026) : une feuille par page (« 1-Neuf - P&P », « 1-F et A Neuf »…), titre en B1,
# « De: » / « À: » en ligne 2, le groupe « Est A (850+) (16 / 16) » en colonne B des
# en-têtes de section, libellés en colonne A (postes de tête en gras, sous-lignes
# non grasses), colonnes Moy. MAD, Moy. MMAP, % d'écart MAD, Moy. PAD, MPAD moy.,
# % d'écart PAD. Le mémo stats n'a que Moy. MAD (et l'écart).
XL_GROUP_RE = re.compile(r"^(.*?)\s*\((\d+)\s*/\s*(\d+)\)\s*$")


def _xl_num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("\u00a0", "").replace(" ", "")
    if not t or t.upper() == "N/A":
        return None
    try:
        return float(t.replace(",", "."))
    except ValueError:
        return None


def _xl_label(v):
    """Libellé au format des PDF : « Neuf - Ventes » → « Neuf ‐ Ventes »."""
    t = re.sub(r"\s+", " ", str(v).replace("\u00a0", " ")).strip()
    return re.sub(r" - ", " \u2010 ", t)


def _xl_head(ws):
    """(titre, début (a, m), fin (a, m), groupe, n, n_total) d'une feuille."""
    import datetime as dt
    title = ws.cell(row=1, column=2).value or ws.title
    debut = fin = None
    grp = n = n_tot = None
    for r in range(1, min(ws.max_row, 12) + 1):
        a, b = ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value
        if isinstance(a, str) and a.strip().lower().startswith("de"):
            d0, d1 = ws.cell(row=r, column=3).value, ws.cell(row=r, column=5).value
            if isinstance(d0, dt.datetime):
                debut = (d0.year, d0.month)
            if isinstance(d1, dt.datetime):
                fin = (d1.year, d1.month)
        if isinstance(a, str) and grp is None:
            m = XL_GROUP_RE.match(a.strip())
            if m:
                grp, n, n_tot = m.group(1), int(m.group(2)), int(m.group(3))
    return _xl_label(title), debut, fin, grp, n, n_tot


def _xl_lignes(ws):
    """[(libellé, gras, [valeurs B…G], en-tête de section)] après les lignes de titre."""
    out = []
    for r in range(4, ws.max_row + 1):
        c = ws.cell(row=r, column=1)
        if c.value is None:
            continue
        b = ws.cell(row=r, column=2).value
        section = isinstance(b, str) and XL_GROUP_RE.match(b.strip()) is not None
        if norm(str(c.value)) == "description":
            continue
        vals = [_xl_num(ws.cell(row=r, column=k).value) for k in range(2, 8)]
        out.append((str(c.value), bool(c.font and c.font.b), vals, section))
    return out


def read_hyundai_xlsx(path):
    """Export Excel d'un rapport eComposite Hyundai → même JSON que la lecture du
    PDF (hyundai_pp, hyundai_ar) ou, pour le mémo stats, {"kind": "hyundai_memo",
    "sections": {feuille: {section: {libellé: Moy. MAD}}}}."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        names = [norm(re.sub(r"^\d+-", "", n)) for n in wb.sheetnames]
        if any(n.endswith("p&p") for n in names):
            kind = "hyundai_pp"
        elif any(n.startswith("inventaire neufs") or n.startswith("employes") for n in names):
            kind = "hyundai_memo"
        elif any(n.startswith("neuf - ventes") or n.startswith("service - vente") for n in names):
            kind = "hyundai_ar"
        else:
            raise ValueError("classeur Excel non reconnu (ni P&P, ni analyse des revenus, ni mémo stats de l'eComposite)")
        pages, head = {}, None
        for ws in wb.worksheets:
            title, debut, fin, grp, n, n_tot = _xl_head(ws)
            if head is None and fin:
                head = (debut, fin, grp, n, n_tot)
            rows = _xl_lignes(ws)
            if kind == "hyundai_pp":
                t = re.sub(r"\s*[-\u2010]?\s*p&p$", "", norm(title)).strip()
                dept = HY_PP_PAGES.get(t)
                if not dept:
                    continue
                bloc, out = None, []
                for lab, _, v, section in rows:
                    nl = norm(lab)
                    b = next((x for k, x in HY_PP_BLOCKS.items() if nl.startswith(k)), None)
                    if section or b:
                        bloc = b or bloc
                        if section:
                            continue
                    if v[0] is None and v[3] is None:
                        continue
                    out.append({"bloc": bloc, "label": _xl_label(lab), "unit": "%" if nl.startswith("%") else "$",
                                "mad": v[0], "mmap": v[1], "pad": v[3], "mpad": v[4]})
                pages[dept] = out
            elif kind == "hyundai_ar":
                items, cur = [], None
                for lab, gras, v, section in rows:
                    q = {"mad": v[0], "mmap": v[1], "pad": v[3], "mpad": v[4]}
                    has = any(x is not None for x in (v[0], v[1], v[3], v[4]))
                    if section or (gras and not has):
                        cur = {"label": _xl_label(lab), "tete": None, "sous": {}}
                        items.append(cur)
                    elif cur is None or (gras and not AR_SUB_RE.match(norm(lab))) or \
                            (not gras and cur["tete"] is None and not cur["sous"] and not AR_SUB_RE.match(norm(lab))):
                        # « Ventes », « PB », « $PB par contrat »… sont des sous-lignes même en gras (total du
                        # bureau commercial) ; sous une simple en-tête de section (« Sommaire » des autres
                        # revenus), une ligne non grasse est un poste (« Retenue », « Incitatif HACC »)
                        cur = {"label": _xl_label(lab), "tete": q, "sous": {}}
                        items.append(cur)
                        if not gras:
                            cur["_simple"] = True
                    elif cur.get("_simple") and not AR_SUB_RE.match(norm(lab)):
                        cur = {"label": _xl_label(lab), "tete": q, "sous": {}, "_simple": True}
                        items.append(cur)
                    else:
                        cur["sous"][_xl_label(lab)] = q
                for it in items:
                    it.pop("_simple", None)
                pages[title] = [it for it in items if it["tete"] or it["sous"]]
            else:
                sec, cur = {}, None
                for lab, _, v, section in rows:
                    if section:
                        cur = sec.setdefault(_xl_label(lab), {})
                        continue
                    if cur is not None and v[0] is not None:
                        cur[_xl_label(lab)] = v[0]
                pages[title] = sec
        if not head or not any(pages.values()):
            raise ValueError("eComposite Excel : période ou chiffres introuvables")
        (y0, m0), (y1, m1), grp, n, n_tot = head
        src = {"hyundai_pp": "P&P", "hyundai_ar": "Analyse des revenus", "hyundai_memo": "Mémo stats"}[kind]
        d = {"kind": kind, "source": f"eComposite Hyundai Canada — {src} (Excel)", "period": f"{y1}-{m1:02d}",
             "debut": f"{y0}-{m0:02d}", "groupe": grp, "n": n, "n_total": n_tot}
        d["sections" if kind == "hyundai_memo" else "pages"] = pages
        return d
    finally:
        wb.close()


# ------------------------------------------------------- bulletin VW Canada
VW_DEALERS = {"volkswagen brossard": "vw"}


def _vw_period(text):
    m = re.search(r"Period:\s*([A-Za-z]{3})[a-z]*\s+(20\d\d)", text)
    if not m:
        return None
    return f"{int(m.group(2))}-{MONTHS_EN[m.group(1).lower()[:3]]:02d}"


def _vw_val(t):
    t = t.strip()
    if t in ("N/A", "", "-"):
        return None
    unit = "%" if t.endswith("%") else ("$" if "$" in t else "")
    x = t.replace("$", "").replace("%", "").replace(",", "")
    try:
        v = float(x)
    except ValueError:
        return None
    return v / 100 if unit == "%" else v


# Indicateurs du « Dealer Performance Summary » (page 4) : (id, début de l'étiquette anglaise)
VW_SUMMARY = [
    ("nv_bpv", "new vehicle sales vs bpv"), ("nv_aged", "new vehicle aged inventory"),
    ("ms_total", "total market share"), ("ms_car", "car market share"), ("ms_truck", "truck market share"),
    ("reg_eff", "registration effectiveness"), ("sales_eff", "sales effectiveness"),
    ("cpo_bpv", "cpo vehicle sales vs bpv"), ("parts_obj", "parts purchased (excluding accessories)"),
    ("parts_loyalty", "parts loyalty"), ("parts_turns", "parts inventory turns"),
    ("acc_pnvr", "accessories $ purch pnvr"), ("ppm", "ppm %"), ("mlc", "maintenance lead conversions"),
    ("cem_sales", "cem sales - average top 3"), ("cem_service", "cem service - average top 3"),
]
# Pages « comparateur » : Service, Parts, New Vehicle, Pre-Owned (id, étiquette)
VW_COMPARATORS = {
    "service": [("absorption", "absorption %"), ("cp_hours_ro", "vw cp sold hours / vw cp ro"),
                ("parts_cp_ro", "parts sales $ / cp ro"), ("ppm_new", "ppm % (new)"), ("ppm_cpo", "ppm % (cpo)"),
                ("mlc", "maintenance lead conversions"), ("car_park", "active car park total"),
                ("tech_eff", "tech efficiency"), ("tech_prod", "tech productivity"),
                ("mix_cp", "labour mix - customer pay"), ("mix_warranty", "labour mix - warranty"),
                ("mix_internal", "labour mix - internal"), ("service_cert", "total service certification")],
    "parts": [("parts_loyalty", "parts loyalty"), ("parts_turns", "parts inventory turns"),
              ("parts_obj", "parts purchased (excluding accessories)"), ("parts_cp_ro", "parts sales $ / cp ro"),
              ("acc_pnvr", "accessories $ purch pnvr"), ("retail_vio", "retail $/vio")],
    "new": [("nv_sales", "new vehicle sales ****"), ("nv_bpv", "new vehicle sales vs bpv"),
            ("nv_aged", "new vehicle aged inventory"), ("ms_total", "total market share"),
            ("sales_eff", "sales effectiveness"), ("reg_eff", "registration effectiveness"),
            ("ms_brand", "total industry market share"), ("vwfs_pen", "volkswagen financial services contract"),
            ("lease_loyalty", "volkswagen finance service lease"), ("loan_loyalty", "volkswagen finance service loan"),
            ("fi_pnv", "f&i pnv")],
    "used": [("cpo_bpv", "cpo vehicle sales vs bpv"), ("cpo_sales", "cpo vehicle sales"), ("cpo_new", "cpo : new ratio"),
             ("noncpo_sales", "non cpo used vw sales"), ("fi_puvr", "f&i puvr")],
}
VW_CEM = [("cem_sales_osat", "cem sales - overall satisfaction"), ("cem_sales_top3", "cem sales - average top 3"),
          ("cem_service_osat", "cem service - overall satisfaction"), ("cem_service_top3", "cem service - average top 3"),
          ("cem_cpo_top3", "cem cpo- average top 3")]


VALUE_TOK = re.compile(r"^(N/A|\$?-?[\d,]*\.?\d+%?|-?\$[\d,]+(?:\.\d+)?|to|\d+/\d+|Met|DQ)$")


def split_values(line):
    """(étiquette, [jetons de valeur]) : les valeurs sont la suite de jetons
    chiffrés à la fin de la ligne ; « 7 to 9 » redevient un seul jeton."""
    toks = line.split(" ")
    i = len(toks)
    while i > 0 and VALUE_TOK.match(toks[i - 1]):
        i -= 1
    # « to » ne peut pas commencer les valeurs
    while i < len(toks) and toks[i] == "to":
        i += 1
    vals, j = [], i
    while j < len(toks):
        if j + 2 < len(toks) and toks[j + 1] == "to":
            vals.append(f"{toks[j]} to {toks[j + 2]}")
            j += 3
        else:
            vals.append(toks[j])
            j += 1
    return " ".join(toks[:i]), vals


def read_vw_report_card(path, text=None):
    text = text or pdf_text(path)
    lines = [re.sub(r"\s+", " ", l).strip() for l in text.splitlines() if l.strip()]
    m = re.search(r"Dealer:\s*(\d+)\s*-\s*([^|\n]+?)\s*\|", text)
    name = m.group(2).strip() if m else ""
    period = _vw_period(text)
    out = {"kind": "vw_report_card", "source": "Volkswagen Canada — Dealer Report Card", "period": period,
           "concession": name, "code": m.group(1) if m else None, "dealer": VW_DEALERS.get(norm(name))}

    # Classement (page 2) : « National(149) 49 -16 -17 » puis « 100 51.50 45 25.0 35 11.5 20 15.0 »
    rk = {}
    for i, l in enumerate(lines):
        mm = re.match(r"(National|Geo Local)\s*\((\d+)\)\s+(\d+)\s+(-?\d+)\s+(-?\d+)$", l)
        if mm:
            k = "national" if mm.group(1) == "National" else "geo"
            if k not in rk:
                rk[k] = {"n": int(mm.group(2)), "rang": int(mm.group(3)), "mom": int(mm.group(4)), "yoy": int(mm.group(5))}
            if k == "national" and "points" not in rk and i + 1 < len(lines):
                p = re.findall(r"-?\d+(?:\.\d+)?", lines[i + 1])
                if len(p) >= 8:
                    p = [float(x) for x in p[:8]]
                    rk["points"] = {"total": [p[1], p[0]], "ventes": [p[3], p[2]], "apres_vente": [p[5], p[4]],
                                    "experience": [p[7], p[6]]}
        mm = re.match(r"(?:August|[A-Z][a-z]+ 20\d\d )?\s*(National|Geo Local) \((\d+)\) (\d+) (\d+) (\d+)$", l)
        if mm:
            k = "national_dept" if mm.group(1) == "National" else "geo_dept"
            rk.setdefault(k, {"total": int(mm.group(3)), "ventes": int(mm.group(4)), "apres_vente": int(mm.group(5))})
    out["classement"] = rk

    # Sommaire des indicateurs (page 4) : 2026 YTD, 2025 YTD, mois, 3 mois,
    # National, Geo Local, [points possibles, points obtenus], rang national,
    # rang Québec, [préqualification]
    summ, in_sum = {}, False
    for l in lines:
        n = norm(l)
        if n.startswith("dealer performance summary 20"):
            in_sum = True
            continue
        if not in_sum:
            continue
        if n.startswith("total 100"):
            p = re.findall(r"\d+(?:\.\d+)?", l)
            summ["_total"] = {"points": float(p[1]) if len(p) > 1 else None}
            in_sum = False
            continue
        lab, vals = split_values(l)
        key = next((k for k, pre in VW_SUMMARY if norm(lab).startswith(pre)), None)
        if not key or key in summ:
            continue
        is_rank = lambda v: re.fullmatch(r"\d+/\d+", v) is not None
        ranks = [v for v in vals if is_rank(v)]
        pq = next((v for v in vals if v in ("Met", "DQ")), None)
        nums = [v for v in vals if not is_rank(v) and v not in ("Met", "DQ")]
        rec = {k: _vw_val(v) for k, v in zip(("ytd", "ytd_ap", "mois", "r3", "national", "geo"), nums[:6])}
        if len(nums) >= 8:
            rec["points_max"], rec["points"] = _vw_val(nums[6]), _vw_val(nums[7])
        if len(ranks) >= 2:
            rec["rang_national"], rec["rang_geo"] = ranks[0], ranks[1]
        if pq:
            rec["prequal"] = "atteint" if pq == "Met" else "non atteint"
        rec["donnees_juillet"] = "*" in lab.replace("***", "")
        summ[key] = rec
    out["sommaire"] = summ

    # Comparateurs (pages 4 à 7) : objectif, 2026 YTD, 2025 YTD, National, Geo Local
    comp, section = {}, None
    heads = {"new vehicle comparator": "new", "pre-owned comparator": "used", "service comparator": "service",
             "parts comparator": "parts"}
    for l in lines:
        n = norm(l)
        h = next((v for k, v in heads.items() if n.startswith(k)), None)
        if h:
            section = h
            continue
        if not section or n.startswith("printed on") or n.startswith("new vehicle models"):
            section = None if section and (n.startswith("printed on") or n.startswith("new vehicle models")) else section
            continue
        lab, vals = split_values(l)
        nl = norm(lab)
        key = None
        for k, pre in VW_COMPARATORS[section]:
            if nl.startswith(pre) and k not in comp.get(section, {}):
                if k == "cpo_sales" and nl.startswith("cpo vehicle sales vs"):
                    continue
                key = k
                break
        if not key or len(vals) < 5:
            continue
        obj, a, b, nat, geo = vals[-5:]
        comp.setdefault(section, {})[key] = {
            "objectif": obj if " to " in obj else _vw_val(obj), "ytd": _vw_val(a), "ytd_ap": _vw_val(b),
            "national": _vw_val(nat), "geo": _vw_val(geo)}
    out["comparateurs"] = comp

    # Expérience client (page 8) : cible, 2026 YTD, 2025 YTD, 3 mois, National, Geo Local
    cem, in_cem = {}, False
    for l in lines:
        if norm(l).startswith("customer experience annual target"):
            in_cem = True
            continue
        if not in_cem:
            continue
        if norm(l).startswith("printed on"):
            break
        lab, vals = split_values(l)
        nl = norm(lab)
        key = next((k for k, pre in VW_CEM if nl.startswith(pre)), None)
        if not key or key in cem:
            continue
        if len(vals) >= 6:
            cem[key] = {k: _vw_val(v) for k, v in zip(("cible", "ytd", "ytd_ap", "r3", "national", "geo"), vals[-6:])}
        elif len(vals) == 5:
            cem[key] = {k: _vw_val(v) for k, v in zip(("ytd", "ytd_ap", "r3", "national", "geo"), vals)}
    out["experience"] = cem
    if not period or len(summ) < 10:
        raise ValueError("bulletin VW : période ou sommaire des indicateurs introuvable")
    return out


# ------------------------------------------ états financiers VW Canada (Excel)
# Un seul gabarit pour l'état de la concession (VWBROSSARD092026.xlsm : en
# anglais depuis 2026, en français avant) et pour les « états financiers
# standard » du groupe de volume (« Financial Statement_Volume Group_A - ≥675_
# Standard statement_AAAAMM_….xlsx » : concessionnaire moyen du groupe A).
# Les chiffres sont repérés par le numéro de ligne officiel (pages 2 à 6), le
# même en 2025 et en 2026, et les colonnes par leurs en-têtes (anglais ou
# français). Demande de Maxime Allard (23 octobre 2026) : comparer VW Brossard
# au groupe A comme dans le rapport de Mustafa Bhaloo.
VW_P2_LIGNES = (1, 2, 4, 7, 15, 27, 47, 62, 63, 64, 73, 74, 75, 76, 80, 81)
VW_P3_LIGNES = (1, 2, 4, 15, 27, 47, 62, 63, 64, 73)
VW_P4_LIGNES = (21, 22, 27, 28, 29, 31, 32, 34, 35, 37, 38, 40, 45, 46, 50)
VW_P5_LIGNES = (1, 3, 4, 13, 18, 22, 29, 41, 52, 53, 54, 62, 63, 64)
VW_FEUILLES = ("Page 1", "Page 2", "Page 3", "Page 4", "Page 5", "Page 6")
VW_MOIS_CLES = ("currentmonth", "moisencours", "moiscourant")
VW_CUMUL_CLES = ("yeartodate", "cumulannuel")


def _vw_cle(v):
    """« CURRENT MONTH » → « currentmonth », « N° LIGNE » → « nligne »."""
    return re.sub(r"[^a-z0-9%&]", "", norm(v)) if isinstance(v, str) else ""


def _vw_cells(ws):
    """{ligne Excel: {colonne: valeur}} (cellules non vides), lu une seule fois."""
    out = {}
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        d = {c: v for c, v in enumerate(row, start=1) if v not in (None, "")}
        if d:
            out[r] = d
    return out


def _vw_nombre(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _vw_lignes(cells, max_r=9):
    """{numéro de ligne officiel: ligne Excel}, d'après la colonne « Line No. » la plus à gauche."""
    cols = [c for r in range(1, max_r + 1) for c, v in (cells.get(r) or {}).items()
            if _vw_cle(v) in VW_LIGNE_CLES and _vw_cle(v) != "nli"]
    if not cols:
        raise ValueError("colonne des numéros de ligne introuvable")
    lc = min(cols)
    out = {}
    for r, d in sorted(cells.items()):
        v = d.get(lc)
        if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer():
            out.setdefault(int(v), r)
    return out


VW_LIGNE_CLES = ("line", "lineno", "ligne", "nligne", "noligne", "nli")


def _vw_blocs(cells):
    """Pages 2 et 3 : [(colonnes du mois, colonnes du cumul)] de gauche à droite (un bloc par département)."""
    # colonnes « Line No. » (à gauche et à droite du tableau) : jamais lues comme des montants
    bornes = sorted({c for r in range(1, 10) for c, v in (cells.get(r) or {}).items() if _vw_cle(v) in VW_LIGNE_CLES})
    for r in range(1, 10):
        d = cells.get(r) or {}
        heads = sorted((c, k) for c, v in d.items() for k in [_vw_cle(v)]
                       if k in VW_MOIS_CLES or k in VW_CUMUL_CLES or k == "%")
        if sum(1 for _, k in heads if k in VW_MOIS_CLES) < 2:
            continue
        def plage(c, meme=()):
            # jusqu'au prochain en-tête ; l'état 2025 répète « CUMUL ANNUEL » sur deux colonnes (pièces) :
            # pour le cumul, on lit les deux
            nxt = min([x for x, kk in heads if x > c and kk not in meme] + [x for x in bornes if x > c] + [c + 9])
            return list(range(c, nxt))
        blocs = []
        for c, k in heads:
            if k in VW_MOIS_CLES:
                y = next((x for x, kk in heads if x > c and kk in VW_CUMUL_CLES), None)
                if y:
                    blocs.append((plage(c), plage(y, VW_CUMUL_CLES)))
        return blocs
    raise ValueError("en-têtes « mois / cumul » introuvables")


def _vw_premier(d, cols):
    for c in cols:
        v = _vw_nombre(d.get(c))
        if v is not None:
            return v
    return None


def _vw_colonnes(cells, groupes):
    """Pages 4 et 5 : {clé: (colonne du mois, colonne du cumul)} d'après la ligne d'en-tête (UNITS, SALES…)."""
    for r in range(1, 10):
        d = cells.get(r) or {}
        trouve = {}
        for cle, noms in groupes.items():
            cs = sorted(c for c, v in d.items() if _vw_cle(v) in noms)
            if len(cs) >= 2:
                trouve[cle] = (cs[0], cs[1])
        if len(trouve) == len(groupes):
            return trouve
    raise ValueError("en-têtes des colonnes introuvables (" + ", ".join(groupes) + ")")


def _vw_periode(cells1):
    """Mois de l'état (« STATEMENT PERIOD » de la page 1) : la date la plus tardive, ou « Sep 2026 »."""
    dates = [v for r in range(1, 10) for v in (cells1.get(r) or {}).values() if isinstance(v, _dt.datetime)]
    if dates:
        d = max(dates)
        return f"{d.year}-{d.month:02d}"
    for r in range(1, 10):
        for v in (cells1.get(r) or {}).values():
            if not isinstance(v, str):
                continue
            m = re.match(r"^\s*([A-Za-zéû]{3,5})[a-zéû]*\.?\s+(\d{4})\s*$", v)
            if m:
                k = norm(m.group(1))[:4]
                mo = MONTHS_EN.get(k[:3]) or MOIS_FR.get(k) or MOIS_FR.get(k[:3])
                if mo:
                    return f"{m.group(2)}-{mo:02d}"
    raise ValueError("période de l'état VW introuvable (page 1)")


def _vw_groupe(cells1):
    """« Groupe A (≥675) » pour l'état standard d'un groupe de volume, sinon None (état d'une concession)."""
    code = nom = None
    for r in range(3, 9):
        for c, v in sorted((cells1.get(r) or {}).items()):
            if isinstance(v, str) and re.fullmatch(r"\s*[A-E]\s*", v):
                code = v.strip()
            elif isinstance(v, str) and re.search(r"[≥>]\s*\d", v):
                nom = re.sub(r"\s+", "", v)
    return f"Groupe {code} ({nom})" if code and nom else (f"Groupe {code}" if code else None)


def vw_statement_lines(path):
    """État financier VW Canada (concession ou groupe de volume) → chiffres du mois et du cumul :
    p2 {ligne: [total, neuf, occasion]}, p3 {ligne: [service, carrosserie, pièces]},
    p4 {ligne: [unités, ventes, PB]}, p5 {ligne: [BR, ventes, PB]}, p6 (heures),
    plus l'inventaire (unités en stock), les techniciens et les employés aux pièces."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if not all(n in wb.sheetnames for n in VW_FEUILLES):
            raise ValueError("pas un état financier VW Canada (pages 1 à 6)")
        P = {n: _vw_cells(wb[n]) for n in VW_FEUILLES}
    finally:
        wb.close()
    out = {"period": _vw_periode(P["Page 1"]), "groupe": _vw_groupe(P["Page 1"]), "month": {}, "ytd": {}}

    for page, lignes, cle in (("Page 2", VW_P2_LIGNES, "p2"), ("Page 3", VW_P3_LIGNES, "p3")):
        cells = P[page]
        rows, blocs = _vw_lignes(cells), _vw_blocs(cells)
        if len(blocs) < 3:
            raise ValueError(f"{page} : trois départements attendus, {len(blocs)} trouvés")
        for i, mode in enumerate(("month", "ytd")):
            out[mode][cle] = {str(n): [_vw_premier(cells.get(rows[n]) or {}, b[i]) for b in blocs[:3]]
                              for n in lignes if n in rows}
    for page, lignes, cle, groupes in (
            ("Page 4", VW_P4_LIGNES, "p4", {"u": ("units", "unites"), "v": ("sales", "ventes"), "pb": ("grossprofit", "profitbrut")}),
            ("Page 5", VW_P5_LIGNES, "p5", {"u": ("ros", "br"), "v": ("sales", "ventes"), "pb": ("grossprofit", "profitbrut")})):
        cells = P[page]
        rows, cols = _vw_lignes(cells), _vw_colonnes(cells, groupes)
        for i, mode in enumerate(("month", "ytd")):
            out[mode][cle] = {str(n): [_vw_nombre((cells.get(rows[n]) or {}).get(cols[k][i])) for k in ("u", "v", "pb")]
                              for n in lignes if n in rows}

    # Page 6 : heures (service), techniciens et employés
    c6 = P["Page 6"]
    rows6 = _vw_lignes(c6)
    d4 = c6.get(rows6.get(4)) or {}
    hc = {"month": [], "ytd": []}
    for c, v in sorted(d4.items()):
        k = _vw_cle(v)
        if "mtd" in k or "mois" in k:
            hc["month"].append(c)
        elif "ytd" in k or "cumul" in k:
            hc["ytd"].append(c)
    d5 = c6.get(rows6.get(5)) or {}
    for mode in ("month", "ytd"):
        # heures vendues client VW : moteur à combustion (1re colonne), puis électriques (2e)
        cs = hc[mode] + [None, None]
        out[mode]["p6"] = {"heures_client": _vw_nombre(d5.get(cs[0])), "heures_client_bev": _vw_nombre(d5.get(cs[1]))}
    for n, cle in ((13, "heures_dispo"), (14, "heures_prod"), (15, "heures_vendues"), (16, "heures_om")):
        d = c6.get(rows6.get(n)) or {}
        txt = min([c for c, v in d.items() if isinstance(v, str) and c > 3] or [10 ** 6])
        nums = [_vw_nombre(v) for c, v in sorted(d.items()) if 2 < c < txt and _vw_nombre(v) is not None]
        for i, mode in enumerate(("month", "ytd")):
            out[mode]["p6"][cle] = nums[i] if len(nums) > i else None
    hdr = c6.get(rows6.get(3)) or {}
    col_serv = next((c for c, v in hdr.items() if _vw_cle(v) == "service"), None)
    col_pa = next((c for c, v in hdr.items() if _vw_cle(v) == "p&a"), None)
    lab = next((c for r in range(1, 12) for c, v in (c6.get(r) or {}).items()
                if _vw_cle(v).startswith(("totalpersonnel", "nombretotal"))), None)
    for r, d in c6.items():
        k = _vw_cle(d.get(lab)) if lab else ""
        if k in ("vwtechnicians", "techniciensvw") and col_serv:
            out["techniciens"] = _vw_nombre(d.get(col_serv))
        if k == "total" and col_pa:
            out["employes_pieces"] = _vw_nombre(d.get(col_pa))

    # Page 1 : unités en stock (neufs VW, occasion au total)
    c1 = P["Page 1"]
    ucol = None
    for r, d in sorted(c1.items()):
        items = sorted(d.items())
        ks = [_vw_cle(v) for c, v in items if isinstance(v, str)]
        if any(k.startswith(("newvehiclesvw1402", "vehiculesneufsvw1402")) for k in ks):
            u = next((c for c, v in items if _vw_cle(v) in ("units", "unites")), None)
            ucol = next((c for c, v in items if u and c > u and _vw_nombre(v) is not None), None)
            out["stock_neuf"] = _vw_nombre(d.get(ucol)) if ucol else None
        elif ucol and any(k.startswith(("totalused", "totaldoccasion")) for k in ks):
            out["stock_occ"] = _vw_nombre(d.get(ucol))
    return out


def read_vw_groupe_xlsx(path):
    """États financiers standard d'un groupe de volume VW (concessionnaire moyen du groupe)."""
    lignes = vw_statement_lines(path)
    if not lignes.get("groupe"):
        raise ValueError("état VW d'une concession, pas d'un groupe de volume")
    return {"kind": "vw_groupe", "period": lignes["period"], "groupe": lignes["groupe"], "lignes": lignes}


# ------------------------------------------------------------- détection
def detect(path):
    """(type, extrait) du PDF, ou (None, None)."""
    try:
        head = pdf_text(path, max_pages=2)
    except Exception:
        return None, None
    h = norm(head)
    if "dealer report card" in h and "volkswagen" in h:
        return "vw_report_card", None
    if ("ef - sommaire" in h or "ef  - sommaire" in h) and ("moy. mad" in h or "moy. pad" in h):
        return "hyundai_ef", None
    if "p&p" in h and "moy. mmap" in h:
        return "hyundai_pp", None
    if "analyse des revenus" in h and "moy. mmap" in h:
        return "hyundai_ar", None
    return None, None


def read_any(path, dealer=None):
    """Lit un composite (PDF, ou export Excel de l'eComposite Hyundai). dealer :
    concession du dossier Drive d'où vient le fichier (le composite Hyundai ne
    nomme pas la concession)."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        feuilles = set(wb.sheetnames)
        wb.close()
        if set(VW_FEUILLES) <= feuilles:
            d = read_vw_groupe_xlsx(path)
            d["dealer"] = dealer or "vw"
            return d
        d = read_hyundai_xlsx(path)
        d["dealer"] = dealer or "hyundai"
        return d
    kind, _ = detect(path)
    if kind == "vw_report_card":
        d = read_vw_report_card(path)
        d["dealer"] = d.get("dealer") or dealer or "vw"
        return d
    if kind in ("hyundai_ef", "hyundai_pp", "hyundai_ar"):
        d = {"hyundai_ef": read_hyundai_ef, "hyundai_pp": read_hyundai_pp, "hyundai_ar": read_hyundai_ar}[kind](path)
        d["dealer"] = dealer or "hyundai"
        return d
    raise ValueError("PDF non reconnu (ni eComposite Hyundai, ni bulletin VW)")


def json_name(d):
    return f"{d['dealer']}_{d['period']}_{d['kind']}.json"


def rounded(x):
    """Arrondit les flottants (0.10800000000000001 → 0.108)."""
    if isinstance(x, float):
        r = round(x, 6)
        return int(r) if r == int(r) and abs(r) >= 1 else r
    if isinstance(x, dict):
        return {k: rounded(v) for k, v in x.items()}
    if isinstance(x, list):
        return [rounded(v) for v in x]
    return x


def write_json(d, outdir):
    """Écrit les chiffres lus (jamais le PDF) dans outdir/<concession>_<mois>_<type>.json."""
    os.makedirs(outdir, exist_ok=True)
    p = os.path.join(outdir, json_name(d))
    with open(p, "w", encoding="utf-8") as f:
        json.dump(rounded(d), f, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
        f.write("\n")
    return p


def load_dir(path):
    """{concession: {mois: {type: données}}} à partir des JSON de path."""
    out = {}
    if not os.path.isdir(path):
        return out
    for name in sorted(os.listdir(path)):
        if not name.endswith(".json"):
            continue
        try:
            with open(os.path.join(path, name), encoding="utf-8") as f:
                d = json.load(f)
        except (OSError, ValueError):
            continue
        dk, pk, kind = d.get("dealer"), d.get("period"), d.get("kind")
        if dk and pk and kind:
            out.setdefault(dk, {}).setdefault(pk, {})[kind] = d
    return out


# ----------------------------------------- postes P&P (composite ↔ état FFS)
# (identifiant, bloc, libellé court, débuts d'étiquette du composite, débuts
# d'étiquette de l'état FFS). Les deux suivent le même ordre de lignes (le
# composite est construit à partir des états FFS des concessions du groupe).
PP_LINES = [
    ("sal_directeurs", "vente", "Salaires des directeurs des ventes", ("salaires - directeurs",), ("salaires - gerants",)),
    ("sal_comm", "vente", "Salaires et commissions — ventes", ("salaires et commissions",), ("salaires et commissions - personnel",)),
    ("sal_fi", "vente", "Salaires et commissions — F&I", ("salaires et commissions f&a",), ("salaires et commissions - bureau",)),
    ("livraison", "vente", "Livraison (préparation à la route)", ("livraison",), ("preparation a la route",)),
    ("pratiques_garantie", "vente", "Pratiques de garantie (travail gratuit)", ("travail a titre gratuit", "pratiques de garantie",
                                                                           "pratique de garantie"), ("pratiques de garantie",)),
    ("publicite", "vente", "Publicité — véhicules", ("publicite et promotion de ventes",), ("publicite et promotion des ventes",)),
    ("coop", "vente", "Rabais de publicité coop", ("rabais de publicite coop",), ("coop rabais publicite",)),
    ("demo", "vente", "Démonstrateurs", ("demonstration",), ("demonstrateurs",)),
    ("entretien_inventaire", "vente", "Entretien de l'inventaire", ("entretien d'inventaire",), ("entretien des vehicules",)),
    ("financement", "vente", "Frais de financement (plan de financement)", ("frais de financement",), ("frais de financement",)),
    ("total_vente", "vente", "Total des frais de vente des véhicules", ("depenses totales de vehicules", "total des frais de vente"),
     ("total des frais de vente des vehicules",)),
    ("comp_gerants", "directs", "Rémunération des gérants", ("compensation - gerants",), ("compensation des directeurs",)),
    ("comp_autre", "directs", "Autre rémunération", ("autre compensation", "compensation - autre"), ("autre compensation",)),
    ("absenteisme", "directs", "Absentéisme et congés payés", ("absenteisme", "remuneration d'absenteisme"), ("absenteisme",)),
    ("accidents", "directs", "Assurance accidents du travail", ("assurance des accidentes",), ("assurance contre les accidents",)),
    ("ae_rpc", "directs", "AE / RRQ", ("ae / rpc",), ("ae / rpc",)),
    ("ass_groupe", "directs", "Assurance groupe et retraite", ("assurance groupe",), ("assurance groupe", "assurance - groupe")),
    ("avantages", "directs", "Avantages sociaux", ("avantages sociaux",), ("avantages sociaux",)),
    ("emploi", "directs", "Sous-total des frais d'emploi", ("sous-total des depenses d'emploi", "sous-total frais d'emploi"),
     ("sous - total frais d'emploi", "sous-total frais d'emploi")),
    ("formation", "directs", "Formation", ("entrainement", "formation"), ("formation",)),
    ("fournitures", "directs", "Fournitures de bureau", ("fournitures de bureau",), ("papeterie",)),
    ("outils", "directs", "Outils d'atelier et fournitures", ("outils",), ("outils",)),
    ("nettoyage", "directs", "Nettoyage et uniformes", ("nettoyage et uniformes",), ("nettoyage",)),
    ("conciergerie", "directs", "Conciergerie et entretien", ("service de concierge",), ("service de conciergerie",)),
    ("ajustements", "directs", "Politiques d'ajustement pièces et service", ("travail a titre gratuit - pieces",),
     ("politiques d'ajustement",)),
    ("pub_ps", "directs", "Publicité — pièces et service", ("publicite et promotion des ventes - pieces",
                                                           "publicite et promotion de ventes - pieces"), ("publicite et promotion des ventes -",)),
    ("entretien_equipement", "directs", "Entretien de l'équipement", ("entretien d'equipement",), ("entretien de l'equipement",)),
    ("vehicules_entreprise", "directs", "Véhicules de l'entreprise", ("vehicules d'entreprise", "depenses - atelier de carrosserie"),
     ("vehicules de l'entreprise",)),
    ("location_equipement", "directs", "Location d'équipement", ("location d'equipement",), ("location d'equipement",)),
    ("logiciel", "directs", "Soutien des logiciels", ("frais de soutien du logiciel",), ("frais de support", "frais support")),
    ("voyages", "directs", "Voyages et repas d'affaires", ("voyages",), ("voyages",)),
    ("telephone", "directs", "Téléphone et Internet", ("telephone",), ("telephone",)),
    ("transport", "directs", "Transport et messagerie", ("transport", "frais de ports"), ("frais postaux",)),
    ("divers", "directs", "Divers", ("divers",), ("divers",)),
    ("total_directs", "directs", "Total des frais directs", ("total des depenses directes", "total des frais directs"),
     ("total des frais directs",)),
    ("loyer", "fixes", "Loyer ou intérêt hypothécaire", ("loyer et/ou interet",), ("interet sur loyer",)),
    ("amort_ameliorations", "fixes", "Amortissement des améliorations locatives", ("amortissement - ameliorations",),
     ("amortissements des ameliorations",)),
    ("entretien_propriete", "fixes", "Entretien de la propriété", ("entretien de propriete",), ("entretien de la propriete",)),
    ("impots_fonciers", "fixes", "Impôts fonciers", ("impots fonciers",), ("impots fonciers",)),
    ("assurance_immeubles", "fixes", "Assurance des bâtiments", ("assurance - immeubles",), ("assurance - batiments",)),
    ("amort_immeubles", "fixes", "Amortissement des immeubles", ("depreciation - batiments",), ("amortissement des immeubles",)),
    ("facteur_location", "fixes", "Sous-total facteur location", ("sous-total facteur de location",), ("sous-total facteur location",)),
    ("taxes", "fixes", "Taxes d'affaires et autres", ("impots autres", "impots d'affaires", "taxes d'affaires"), ("taxes d'affaires",)),
    ("amort_equipement", "fixes", "Amortissement de l'équipement", ("depreciation - equipement",), ("amortissement - equipement",)),
    ("assurance_generale", "fixes", "Assurance générale", ("assurance generale",), ("assurances generales",)),
    ("chauffage", "fixes", "Chauffage, électricité et eau", ("chauffage",), ("chauffage",)),
    ("honoraires", "fixes", "Honoraires professionnels", ("honoraires",), ("frais professionnels",)),
    ("adhesions", "fixes", "Adhésions et abonnements", ("affiliations",), ("frais d'adhesion",)),
    ("dons", "fixes", "Dons", ("des dons",), ("dons",)),
    ("total_indirects", "fixes", "Total des frais indirects", ("total des depenses indirectes", "total des frais indirects"),
     ("total des frais indirects",)),
    ("total_frais", "fixes", "Total des dépenses", ("total des frais", "total des depenses"), ("total des depenses",)),
]
PP_BY_ID = {x[0]: x for x in PP_LINES}
# Blocs de l'état FFS : ligne d'en-tête → ligne de total
FFS_BLOCKS = {"vente": ("frais de vente des vehicules", "total des frais de vente des vehicules"),
              "directs": ("frais directs", "total des frais directs"),
              "fixes": ("frais indirects", "total des depenses")}


def pp_canon(rows):
    """Lignes P&P du composite d'un département → {id: ligne}."""
    out = {}
    for r in rows:
        lab = norm(r["label"]).replace("‐", "-")
        lab = re.sub(r"\s*-\s*", " - ", lab)
        for key, bloc, _, cpre, _ in PP_LINES:
            if key in out:
                continue
            if bloc == "vente" and r.get("bloc") not in ("vente", None):
                continue
            if bloc != "vente" and r.get("bloc") == "vente" and key not in ("total_vente",):
                continue
            pres = [re.sub(r"\s*-\s*", " - ", p) for p in cpre]
            if any(lab.startswith(p) for p in pres):
                # « Salaires et commissions » ne doit pas prendre la ligne F&A ; « Total des frais »
                # (total) ne doit pas prendre « Total des frais de vente / directs / indirects »
                if key == "sal_comm" and "f&a" in lab:
                    continue
                if key == "pratiques_garantie" and "pieces" in lab:
                    continue
                if key == "publicite" and "pieces" in lab:
                    continue
                if key == "total_frais" and any(w in lab for w in ("vente", "direct", "indirect")):
                    continue
                if key == "vehicules_entreprise" and lab.startswith("depenses - atelier") and "vehicules" not in lab:
                    continue
                out[key] = r
                break
    return out


# ------------------------------------- valeurs de la concession (état FFS)
def ffs_pp_lines(ws, rows, col, num):
    """Postes PP_LINES d'une colonne de l'état FFS (repérés dans leur bloc)."""
    out = {}
    lab_rows = [(r, re.sub(r"\s*-\s*", " - ", lab)) for r, lab in rows]
    bounds = {}
    for bloc, (head, tot) in FFS_BLOCKS.items():
        h = next((r for r, lab in lab_rows if lab == head or lab.startswith(head + " ") and bloc != "vente"
                  or (bloc == "vente" and lab == head)), None)
        t = next((r for r, lab in lab_rows if lab.startswith(re.sub(r"\s*-\s*", " - ", tot)) and (h is None or r > h)), None)
        bounds[bloc] = (h, t)
    for key, bloc, _, _, fpre in PP_LINES:
        h, t = bounds[bloc]
        if h is None or t is None:
            continue
        pres = [re.sub(r"\s*-\s*", " - ", p) for p in fpre]
        for r, lab in lab_rows:
            if not (h <= r <= t):
                continue
            if key == "publicite" and " - " in lab:
                continue
            if key == "total_frais" and r != t:
                continue
            if any(lab.startswith(p) for p in pres):
                v = num(ws, r, col)
                out[key] = v if v is not None else 0.0
                break
    return out


def hyundai_statement_lines(path):
    """Chiffres de Hyundai Longueuil aux définitions du composite, lus dans
    l'état financier Hyundai Canada (FFS) : {"month"|"ytd": {"ef": {dept:
    {id: v}}, "pp": {dept: {label_normalisé: v}}}}.

    Définitions vérifiées sur les moyennes du composite d'août 2026 :
    profit net avant bonis = profit net avant impôt + salaire et bonis des
    propriétaires ; absorption = PB pièces + service + carrosserie ÷ (frais
    total − frais de vente des véhicules) ; publicité = publicité des
    véhicules + publicité pièces, service ; PVND / PVOD = bureau commercial
    ÷ unités détail."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws2, ws3, ws4 = wb["Page 2"], wb["Page 3"], wb["Page 4"]

    def num(ws, r, c):
        if r is None:
            return None
        v = ws.cell(row=r, column=c).value
        if not isinstance(v, (int, float)) or isinstance(v, bool):
            return None
        left = ws.cell(row=r, column=c - 1).value
        return -v if isinstance(left, str) and left.strip() == "(" else float(v)

    def rows_of(ws, col):
        out = []
        for r in range(1, ws.max_row + 1):
            lab = ws.cell(row=r, column=col).value
            if isinstance(lab, str) and lab.strip():
                out.append((r, norm(lab)))
        return out

    r2, r3, r4 = rows_of(ws2, 2), rows_of(ws3, 4), rows_of(ws4, 34)

    def find(rows, prefix, start=0):
        p = norm(prefix)
        for r, lab in rows:
            if r >= start and lab.startswith(p):
                return r
        return None

    cols = {"month": {"total": (ws2, 25), "neuf": (ws2, 51), "occasion": (ws2, 69), "pieces": (ws3, 20),
                      "service": (ws3, 38), "carrosserie": (ws3, 56), "location": (ws3, 75)},
            "ytd": {"total": (ws2, 38), "neuf": (ws2, 60), "occasion": (ws2, 78), "pieces": (ws3, 29),
                    "service": (ws3, 47), "carrosserie": (ws3, 65), "location": (ws3, 84)}}
    p4 = {"month": {"units": 2, "ventes": 7, "pb": 17}, "ytd": {"units": 53, "ventes": 58, "pb": 68}}
    out = {}
    for mode in ("month", "ytd"):
        ef, pp = {}, {}
        for dept, (ws, c) in cols[mode].items():
            rows = r2 if ws is ws2 else r3
            g = lambda prefix, start=0: num(ws, find(rows, prefix, start), c)
            ventes, pb = g("ventes totales"), g("profit brut total")
            frais, po = g("total des depenses"), g("profit d'operation net")
            e = {"ventes": ventes, "pb": pb, "frais": frais, "profit_op": po,
                 "profit_op_pct_pb": (po / pb) if po is not None and pb else None}
            ef[dept] = e
            pp[dept] = ffs_pp_lines(ws, rows, c, num)
        # Totaux de la concession
        t = ef["total"]
        ws, c = cols[mode]["total"]
        g2 = lambda prefix, start=0: num(ws2, find(r2, prefix, start), c) or 0.0
        r_po = find(r2, "profit d'operation net")
        pn = g2("profit net", r_po or 0)
        sp = g2("salaire - proprietaires", r_po or 0)
        bp = g2("bonis - proprietaires", r_po or 0)
        t["pn_avant_bonis"] = pn + sp + bp
        t["remuneration_proprietaires"] = sp + bp   # profit net avant impôts = pn_avant_bonis − ce montant
        t["pn_pct_pb"] = t["pn_avant_bonis"] / t["pb"] if t["pb"] else None
        pub_v = g2("publicite et promotion des ventes")
        r_dir = find(r2, "frais directs")
        pub_ps = g2("publicite et promotion des ventes", (r_dir or 0) + 1)
        t["publicite"] = pub_v + pub_ps
        fv = g2("total des frais de vente des vehicules")
        fo_pb = sum((ef[d]["pb"] or 0) for d in ("pieces", "service", "carrosserie"))
        t["absorption"] = fo_pb / (t["frais"] - fv) if t["frais"] and (t["frais"] - fv) else None
        # Unités et bureau commercial (Page 4)
        cc = p4[mode]
        u = lambda prefix: num(ws4, find(r4, prefix), cc["units"])
        pbc = lambda prefix: num(ws4, find(r4, prefix), cc["pb"])
        un, uo = u("total hyundai detail - neufs"), u("total ventes au detail")
        fin, fio = pbc("bureau commercial - neuf"), pbc("bureau commercial - total")
        ef["neuf"].update({"unites": un, "fi": fin, "fi_unite": (fin / un) if fin is not None and un else None})
        ef["occasion"].update({"unites": uo, "fi": fio, "fi_unite": (fio / uo) if fio is not None and uo else None})
        # Analyse des revenus : service et pièces par type (Page 4 : nombre, ventes, PB)
        ar = {}
        heads = {"service": next((r for r, lab in r4 if lab == "service"), None),
                 "pieces": next((r for r, lab in r4 if lab == "pieces"), None)}
        for _, key, _, fpre, head in AR_LINES:
            start = heads.get(head)
            if start is None:
                continue
            r = find(r4, fpre, start)
            if r is None:
                continue
            ar[key] = {"n": num(ws4, r, cc["units"]), "ventes": num(ws4, r, cc["ventes"]), "pb": num(ws4, r, cc["pb"])}
        out[mode] = {"ef": ef, "pp": pp, "ar": ar, "x": ffs_extra(wb, r4, cc, mode, num, find)}
    wb.close()
    return out


def ffs_extra(wb, r4, cc, mode, num, find):
    """Lignes de l'état FFS utilisées par l'analyse au composite (format du
    rapport de Mustafa Bhaloo, demande de Maxime Allard du 22 octobre 2026) :
      - Page 4 : unités du département neuf (flotte comprise, sans les échanges
        entre concessionnaires), détail et gros d'occasion ;
      - Page 5 : retenue et incitatifs HACC (autres revenus), total des autres
        revenus et des autres dépenses ;
      - Page 6 : heures de l'atelier, bons de réparation, personnel.
    Vérifié sur l'état de septembre 2026 contre le rapport de Mustafa : 76 unités
    neuves, retenue 79 360 $ + HACC 11 250 $ = 90,6 k$, 16 unités au détail,
    38 en gros, 600 BT client, 10 techniciens, 3 employés aux pièces."""
    ws4 = wb["Page 4"]
    cols4 = {"n": cc["units"], "ventes": cc["ventes"], "pb": cc["pb"]}

    def p4(prefix, champ="pb"):
        return num(ws4, find(r4, prefix), cols4[champ])

    x = {"u_neuf": p4("total depart. veh. neufs", "n"), "u_flotte": p4("flotte national", "n"),
         "u_occ_det": p4("total ventes au detail", "n"), "ventes_occ_det": p4("total ventes au detail", "ventes"),
         "pb_occ_det": p4("total ventes au detail"),
         "u_gros": p4("ventes en gros", "n"), "pb_gros": p4("ventes en gros")}

    def lignes(ws, col):
        out = []
        for r in range(1, ws.max_row + 1):
            lab = ws.cell(row=r, column=col).value
            if isinstance(lab, str) and lab.strip():
                out.append((r, norm(lab)))
        return out

    if "Page 5" in wb.sheetnames:
        ws5 = wb["Page 5"]
        rev, dep = lignes(ws5, 2), lignes(ws5, 42)
        c_rev, c_dep = (25, 66) if mode == "month" else (34, 75)
        v5 = lambda rows, prefix, c: num(ws5, find(rows, prefix), c)
        x.update({"retenue": v5(rev, "retenue", c_rev) or 0.0, "hacc": v5(rev, "incitatif hacc", c_rev) or 0.0,
                  "autres_revenus": v5(rev, "total autres revenus", c_rev), "autres_depenses": v5(dep, "total autres depenses", c_dep)})
        # Inventaire par âge à la fin du mois (1-30, 31-60, 61-90, plus de 90 jours) : unités neuves en
        # colonne 11, d'occasion en colonne 43 (septembre 2026 : 180 et 55 — mois d'inventaire 2,4 et 3,4 chez Mustafa)
        ages = [r for r, lab in rev if re.match(r"^(1 ?- ?30|31 ?- ?60|61 ?- ?90|plus de 90) jours", lab)]
        if len(ages) >= 4:
            x["stock_neuf"] = sum(num(ws5, r, 11) or 0.0 for r in ages[:4])
            x["stock_occ"] = sum(num(ws5, r, 43) or 0.0 for r in ages[:4])
    if "Page 6" in wb.sheetnames:
        ws6 = wb["Page 6"]
        st = lignes(ws6, 49)
        c6 = 60 if mode == "month" else 66
        v6 = lambda prefix: num(ws6, find(st, prefix), c6)
        x.update({"heures_dispo": v6("temps disponible"), "heures_pointees": v6("heures poinconnees"),
                  "heures_facturees": v6("temps facture"), "heures_client": v6("temps facture - client"),
                  "br_client": v6("bons de reparation - client"), "br_total": v6("total des bons de reparation")})
        pers = lignes(ws6, 2)
        rt = find(pers, "techniciens")
        x["techniciens"] = num(ws6, rt, 17)
        r_tot = next((r for r, lab in pers if lab == "total" and r > (rt or 0)), None)
        x["employes_pieces"] = num(ws6, r_tot, 48)
    return x


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description="Lecture des composites des constructeurs (PDF → JSON)")
    ap.add_argument("pdf", nargs="+")
    ap.add_argument("--sortie", default=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                                      "sources", "composites"))
    a = ap.parse_args(argv)
    for p in a.pdf:
        d = read_any(p)
        print(f"{os.path.basename(p)} → {write_json(d, a.sortie)}")


if __name__ == "__main__":
    main()
